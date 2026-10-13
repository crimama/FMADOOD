"""Reinforced ScopeAD (Outline v18) deck: detailed method, bank construction and bagging rationale, and results
shown as the outline's own tables (with its bold / underline marks). DSBA template frame; titles are plain noun
phrases and leads are single sentences without a label.
usage: python3 build_deck_v18b.py out.pptx   (figures: charts_v18.py, figs_v2.py; overview = PaperWorks figure)
"""
import os
import re
import sys
from pathlib import Path

sys.path.insert(0, os.path.expanduser('~/.claude/skills/dsba-slide-deck/scripts'))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from deck_style import (BOLD, FONT, INK, LOGO, NAVY, SUB, XBOLD, _textbox, figure, frame_slide,  # noqa: E402
                        new_deck, style_run, table, text_box)
from deck_ops import set_sections  # noqa: E402
from outline_tables import _clean, table_raw  # noqa: E402

CH = Path(__file__).resolve().parent / 'charts'
FIG = Path('/home/hun/Volume/RESEARCH/FSAD-TTA/PaperWorks/figures')
OUT = sys.argv[1]
SRC = '출처: PaperWorks/outlines/Outline_v18.md'
MAPNOTE = '맵은 캐시된 bag 평균 계층별 정규화 맵을 z_d + 0.125·z5로 재구성한 표시용 예시 (새 평가 아님)'

ORDER = ['cover', 'toc', 'summary',
         'bg', 'ambig', 'align', 'p1', 'p2a', 'p2b', 'contrib', 'related',
         'overview', 'steps', 'bankcmp', 'resol', 'bagging', 'whybag', 'scale_eq', 'scale_read', 'control_eq',
         'control_ex', 'aggregate', 'spec',
         't1', 't2', 't3', 't4', 't5', 't5b', 't6', 't7',
         't8', 't8a', 't8b', 't8c', 't8d', 't10', 't10a', 't10b', 't11', 't12',
         'discussion', 'pending', 'conclusion',
         'a_claims']
MAIN_LAST = ORDER.index('conclusion') + 1
P = {k: i + 1 for i, k in enumerate(ORDER)}
A = {k: f'A-{i + 1}' for i, k in enumerate(ORDER[MAIN_LAST:])}


def page(key):
    return A.get(key, P[key])


S1, S2, S3, S4, S5 = ('01. Introduction', '02. Method', '03. 실험 결과', '04. Ablation · 분석', '05. 논의와 향후 계획')
SA = 'Appendix A. 연구 노트'
prs = new_deck()
slides = {}


def em(text):
    return sum(1.0 if ord(c) > 0x1100 else (0.3 if c == ' ' else 0.6) for c in text)


def lint(key, title, lead, detail):
    if ':' in title:
        print(f'WARN colon title  p{page(key)}: {title}')
    if em(title) * 28 / 72 > 10.6:
        print(f'WARN title wraps  p{page(key)}: {title}')
    if em(lead) * 16 / 72 > 11.3:
        print(f'WARN lead wraps   p{page(key)}: {lead} ({em(lead) * 16 / 72:.1f})')
    for line in detail:
        if em(line) * (15 if line.startswith('- ') else 16) / 72 > 11.6:
            print(f'WARN bullet wraps p{page(key)}: {line} ({em(line) * 16 / 72:.1f})')


def slide(key, section, title, lead, detail=(), note='', notes=''):
    lint(key, title, lead, detail)
    s, top = frame_slide(prs, section=section, title=title, lead=lead, detail=detail, note=note, notes=notes,
                         page=page(key))
    slides[key] = s
    return s, top


def boxed(s, x, y, w, h, label, text):
    sh = _textbox(s, x, y, w, h, 'KeyBox')
    sh.line.color.rgb = RGBColor.from_string(NAVY)
    sh.line.width = Pt(1)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(.2)
    tf.margin_top = tf.margin_bottom = Inches(.12)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = label + '  '; style_run(r, 15, NAVY, BOLD, True)
    r = p.add_run(); r.text = text; style_run(r, 15, INK, FONT, False)


def pretty(cell):
    c = _clean(cell).replace('$', '').replace('lambda', 'λ').replace('Delta', 'Δ').replace(',z_5', '·z5')
    c = c.replace('z_5', 'z5').replace('^*', '*').replace('\\', '')
    return re.sub(r'\s+', ' ', c).strip()


def raw_table(s, caption, y, widths, *, header=None, size=10, row_h=None, bottom=6.72, keep=None, x=.70,
              strong_cols=(0,), name='SummaryTable', mod=None, tight=False, plain=True):
    """Outline table as-is; cells marked ** in the outline become bold navy, <u> cells are underlined."""
    raw = table_raw(caption)
    head = header or [pretty(c) for c in raw[0]][:len(widths)]
    body_raw = [r[:len(widths)] for r in raw[1:] if (keep is None or keep(r))]
    body = [[pretty(c) for c in r] + [''] * (len(widths) - len(r)) for r in body_raw]
    if mod:
        body = [mod(r) for r in body]
    frame = table(s, x, y, widths, head, body, size=size, row_h=row_h, bottom=bottom, strong_cols=strong_cols,
                  name=name)
    tb = frame.table
    if tight:  # dense tables: minimal vertical cell padding so rows keep their set height
        for row in tb.rows:
            for cell in row.cells:
                cell.margin_top = cell.margin_bottom = Inches(.005)
    for i, r in enumerate(body_raw, 1):
        for j, cell in enumerate(r):
            if '**' in cell or '<u>' in cell:
                for p in tb.cell(i, j).text_frame.paragraphs:
                    for run in p.runs:
                        if '**' in cell:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor.from_string(NAVY)
                        if '<u>' in cell:
                            run.font.underline = True
    if plain:
        plain_table(frame)
    return frame


def stacked(s, paths, top, bottom=6.72, gap=.1):
    from PIL import Image
    ratios = [Image.open(p).size[1] / Image.open(p).size[0] for p in paths]
    w = min(11.94, (bottom - top - gap * (len(paths) - 1)) / sum(ratios))
    y = top
    for p, r in zip(paths, ratios):
        figure(s, p, y, bottom=y + w * r)
        y += w * r + gap


PLAIN_STYLE = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'  # No Style, No Grid -> no theme fills
LINE_W = 9525  # 0.75 pt


def plain_table(frame):
    """Result-table look: no background, thin black grid lines, black text (bold / underline kept)."""
    from pptx.oxml.ns import qn
    from lxml import etree
    tbl = frame._element.graphic.graphicData.tbl
    tblPr = tbl.tblPr
    tblPr.set('firstRow', '0')
    tblPr.set('bandRow', '0')
    tblPr.find(qn('a:tableStyleId')).text = PLAIN_STYLE
    for tc in tbl.iter(qn('a:tc')):
        tcPr = tc.get_or_add_tcPr()
        for child in list(tcPr):
            if child.tag in {qn(t) for t in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB', 'a:solidFill', 'a:noFill',
                                             'a:gradFill', 'a:blipFill', 'a:pattFill', 'a:grpFill')}:
                tcPr.remove(child)
        for pos, tag in enumerate(('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB')):
            ln = etree.Element(qn(tag), w=str(LINE_W), cap='flat', cmpd='sng', algn='ctr')
            fill = etree.SubElement(ln, qn('a:solidFill'))
            etree.SubElement(fill, qn('a:srgbClr'), val='000000')
            etree.SubElement(ln, qn('a:prstDash'), val='solid')
            tcPr.insert(pos, ln)
        tcPr.insert(4, etree.Element(qn('a:noFill')))
    for row in frame.table.rows:
        for cell in row.cells:
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.color.rgb = RGBColor.from_string('000000')


def table_bottom(y, n_rows, row_h):
    return y + .42 + n_rows * row_h


# ---------------------------------------------------------------- cover / toc / summary
def cover():
    blank = next(l for l in prs.slide_layouts if l.name == '빈 화면')
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(str(LOGO), Inches(.9), Inches(1.15), Inches(1.55), Inches(.81)).name = 'Logo'
    text_box(s, .9, 2.25, 11.5, .8, 'ScopeAD', 44, NAVY, XBOLD, True, name='Title')
    text_box(s, .9, 3.10, 11.8, .9, ['Observation Expansion and Support-Calibrated Fusion',
                                    'for Few-Shot Industrial Anomaly Localization'], 20, NAVY, BOLD, True, name='Subtitle')
    text_box(s, .9, 4.0, 11.5, .5, '소수 정상 이미지에서 관측을 늘리고, 관측별 점수를 정상 기준으로 맞춰 결합하는 결함 국소화', 18,
             INK, FONT, name='Subtitle2')
    line = s.shapes.add_connector(1, Inches(.9), Inches(4.75), Inches(12.4), Inches(4.75))
    line.line.color.rgb = RGBColor.from_string(NAVY)
    line.line.width = Inches(.015)
    text_box(s, .9, 4.95, 11.5, .4, '연구 아웃라인 v18 발표  ·  2026-10-12', 16, SUB, FONT, name='Date')
    s.notes_slide.notes_text_frame.text = ('근거 문서: PaperWorks/outlines/Outline_v18.md. ScopeAD = Support-Calibrated '
                                           'Observation-expanded Patch Ensemble for Anomaly Detection (이전 이름 SCARF).')
    slides['cover'] = s


def toc():
    rng = lambda a, b: f'{P[a]}–{P[b]}쪽'  # noqa: E731
    slide('toc', '목차', '발표 순서', '문제와 주장 → 방법 → 주 결과 → 구성 요소 분석 → 남은 검증 순으로 진행한다',
          [f'01 Introduction: 왜 관측을 늘리고, 왜 관측별로 맞춰 결합하는가 · 기여 ({rng("bg", "related")})',
           f'02 Method: 절차, 메모리 뱅크와 bag, Scale Calibration, Deep-Anchored Weighting, 맵 통합 ({rng("overview", "spec")})',
           f'03 실험 결과: 표 1–7 (설정, few-shot 주 비교, MVTec AD·VisA, private, 분포 변화) ({rng("t1", "t7")})',
           f'04 Ablation · 분석: 표 8–8d, 10–10b, 11, 12 ({rng("t8", "t12")})',
           f'05 논의와 향후 계획: 주장 범위와 남은 검증 ({rng("discussion", "conclusion")})',
           f'Appendix A: Claim–evidence–status matrix ({A["a_claims"]})'])


def summary():
    s, top = slide('summary', '요약', '한 장 요약',
                   '관측을 늘리고 관측별 점수를 맞춰 결합해, AD2 16-shot에서 SuperAD보다 AUPRO@0.05 +4.17pp')
    rows = [
        ['문제', '소수 정상 support에서 관측을 어떻게 늘리고 합칠까?', '큰 거리의 원인이 모호하고, 관측을 늘리면 단위·정상 반응이 섞임',
         f'{P["ambig"]}–{P["p2b"]}'],
        ['방법', '무엇을 하나?', '특징·해상도·bag별 48개 bank + 정상 q95 단위 정렬 + Shallow 고정 가중(λ = 0.125)',
         f'{P["overview"]}–{P["aggregate"]}'],
        ['주 결과', 'AD2 few-shot 경쟁력은?', '1–16 shot 3개 지표 모두 1위, 16-shot 3-seed 60.06 vs 55.89', str(P['t4'])],
        ['구성 분석', '이득은 어디서 오나?', '다해상도 +4.95 / +2.10pp, 단일 bank 대비 +2.09pp, λ = 1 대비 +5.55pp',
         f'{P["t8"]}–{P["t10b"]}'],
        ['한계', '아직 확인 안 된 것은?', 'q95의 정규화 생략 대비 이득 없음, linear0125 private·shift·비용 미확보',
         f'{P["t8c"]}·{P["pending"]}'],
    ]
    table(s, .70, top + .05, [1.4, 3.6, 5.74, 1.2], ['섹션', '질문', '결론', '쪽'], rows, size=13, bottom=5.9)


# ---------------------------------------------------------------- 01 introduction
def bg():
    s, top = slide('bg', S1, '연구 배경',
                   '재학습 없이 정상 1–16장만으로 결함 영역을 찾는 few-shot 이상 국소화를 다룬다',
                   ['동결 백본 재사용: 대규모 사전학습 표현(DINOv2)은 별도 학습 없이 여러 과제로 전이됨',
                    '패치 매칭: query 패치와 가장 가까운 정상 패치의 거리로 결함 위치를 찾음 → 제품별 재학습 불필요',
                    '주 지표 AUPRO@0.05: 오경보율 5% 이하에서 결함 영역을 얼마나 덮는가'],
                   note=f'예시: MVTec AD hazelnut. {MAPNOTE}. {SRC} §1.1 (Oquab 2023; Roth 2022; Damm 2025)')
    figure(s, CH / 'task.png', top + .05)


def ambig():
    s, top = slide('ambig', S1, '문제 정의',
                   '큰 거리는 결함에서도 support에 없던 정상 변동에서도 생겨, 소수 정상 이미지로는 원인을 가리기 어렵다',
                   ['High-Score Ambiguity: "큰 거리"만으로는 결함인지 미관측 정상 변동(조명·자세·반사)인지 알 수 없음',
                    '① 단서 부족: 정상 이미지를 늘릴 수 없어, 미세 결함과 정상 변동을 가를 시각 단서 자체가 적음',
                    '② 결합의 어려움: 관측을 늘리면 관측마다 거리 척도가 다르고(②-a) Shallow의 정상 반응이 섞임(②-b)'],
                   note=f'그림은 개념도(가상 분포)이며 실측이 아님. 일반 vision 근거: 분포 변화 아래 성능 저하(Hendrycks 2019), '
                        f'점수 신뢰도 변화(Ovadia 2019). {SRC} §1.1–1.2')
    figure(s, CH / 'ambiguity_concept.png', top + .05)


def align():
    s, top = slide('align', S1, '문제 · 동기 · 기여의 대응',
                   '각 문제에 동기와 방법 방향을 하나씩 대응시키고, ③에서 두 종류의 이득을 분리 검증한다',
                   note=f'①은 관측 단계, ②-a·②-b는 결합 단계의 문제. 근거 수치는 AD2 test_public seed 0 개발 비교(표 8, 8b, 8d, 10a)와 '
                        f'16-shot 3-seed 주 비교(표 4). 문제별 상세는 {P["p1"]}–{P["p2b"]}쪽. {SRC} §1.1–1.3')
    figure(s, CH / 'alignment.png', top + .05)


def p1():
    s, top = slide('p1', S1, '① 단서 부족 → 관측 확장',
                   '정상 이미지를 늘릴 수 없으면, 같은 이미지를 여러 깊이·해상도에서 관찰해 단서를 늘린다',
                   ['문제: few-shot에서는 정상 표본이 희소해, 미세 결함을 가를 시각 단서가 부족함',
                    '동기: 다계층(DPT, Raghu 2021)·다중 척도(FPN, S², FeatUp) 관찰이 dense prediction을 개선',
                    '방향: 4 해상도(토큰 1개 = 짧은 변 3.1% → 1.25%) × 3 계층(Deep 11·23, Shallow 5)으로 관측 확장',
                    '근거: 4 해상도 결합이 최선 단일 해상도보다 AUPRO@0.05 +4.95pp(2-shot), +2.10pp(16-shot) (표 10a)'],
                   note=f'아래: hazelnut query의 해상도별 ScopeAD 맵. {MAPNOTE}. 높은 해상도는 기존 영상의 세밀한 표본화이며 새 정상 변동을 '
                        f'만들지 않음. {SRC} §1.2 (1)(2), §2.2, 표 10a')
    figure(s, CH / 'resolution_row.png', top + .1)


def p2a():
    s, top = slide('p2a', S1, '②-a 관측 간 척도 차이 → 단위 정렬',
                   '관측마다 bank를 따로 두면 정상 거리 척도가 달라지므로, 결합 전에 정상 기준으로 단위를 맞춘다',
                   ['문제: 계층 특징을 한 bank로 합치면 정상 변동이 섞이고, 나누면 bank마다 거리 척도가 다름 (아래 왼쪽)',
                    '동기: 점수 맵의 정규화가 결과를 바꾸고(Kim 2021), 이질적 점수는 결합 전 척도를 맞춤(ViM 2022)',
                    '방향: 관측별 bank를 유지하고, 각 bank의 정상 LSO 거리 q95를 1로 정렬 (아래 오른쪽)',
                    '근거: 단일 bank 57.48% vs 관측별 bank 59.57% (+2.09pp, 표 8), q95 > MAD +0.76–1.69pp (표 8d)'],
                   note='아래: 실제 support LSO 거리 (AD2 walnuts, 16-shot, 672 px, bag 0, 계층당 4,096개; 캐시). 원고 진단: 12개 컴포넌트의 q95 척도 '
                        f'차이 평균 약 1.8배. q95의 정규화 생략 대비 이득은 주장하지 않음. {SRC} §1.2 "관측 결합의 어려움", 표 8, 8d')
    figure(s, CH / 'lso_scales.png', top + .1)


def p2b():
    s, top = slide('p2b', S1, '②-b Shallow 정상 반응 → Deep 기준 가중',
                   '단위를 맞춰도 Shallow 점수에는 정상 질감 반응이 섞이므로, Deep을 기준선으로 Shallow는 작게 더한다',
                   ['문제: Shallow(L5)는 미세 결함에 민감하지만 정상 질감·외관에도 반응 → λ = 1이면 정상 walnut 전체가 밝아짐',
                    '동기: Deep(L11·23)은 조명·색 변화에 둔감한 안정적 의미 기준선을 제공',
                    '방향: Deep-Anchored Weighting — S = z_d + 0.125·z5 (임계값·학습 가중치 없음)',
                    '근거: Shallow 전체 가산(λ = 1) 54.02% vs λ = 0.125 59.57% (+5.55pp, 표 8b)'],
                   note=f'AD2 walnuts 정상·결함 query, 하늘색 = GT(표시용). {MAPNOTE}. λ는 AD2 2-shot seed 0로 선택(post-hoc). {SRC} §1.2 (1), 표 8b')
    figure(s, CH / 'shallow_leak.png', top + .05)


def contrib():
    s, top = slide('contrib', S1, '주요 기여',
                   '①의 관측 확장과 ②의 보정·가중 결합을 제안하고, ③에서 두 이득을 통제 실험으로 분리해 검증한다',
                   note=f'모두 개발에 사용한 모집단(AD2 public, MVTec AD, VisA) 결과. "SOTA"는 Outline §1.3 표현. {SRC} §1.3')
    rows = [
        ['① Observation\nExpansion', '① 단서 부족', '재학습 없이 해상도(토큰 3.1% → 1.25%)와\n계층(Shallow / Deep) 두 관측 축을 확장',
         '최선 단일 해상도 대비\n+4.95 / +2.10pp', f'{P["t10a"]}·{P["t10"]}'],
        ['② Calibrated &\nDeep-Anchored Fusion', '②-a 척도 차이', '관측별 bank + 정상 q95 척도 보정',
         '단일 bank 대비 +2.09pp,\nMAD 대비 +0.76–1.69pp', f'{P["t8"]}·{P["t8d"]}'],
        ['', '②-b Shallow 정상 반응', 'Deep 기준선 + Shallow 고정 가중 가산 (λ = 0.125),\n학습 없는 결정론적 파이프라인',
         'Shallow 전체 가산 대비\n+5.55pp', str(P['t8b'])],
        ['③ Disentanglement\n& SOTA', '①·② 검증', '관측 확장(구성) 이득과 결합 이득을\n통제 실험으로 분리, AD2 few-shot 비교',
         'SuperAD 대비 3-seed\nAUPRO@0.05 +4.17pp', str(P['t4'])],
    ]
    frame = table(s, .70, top + .15, [2.5, 2.0, 4.0, 2.54, .9], ['기여', '해결하는 문제', '방법', '근거', '쪽'], rows, size=12,
                  bottom=6.0)
    frame.table.cell(2, 0).merge(frame.table.cell(3, 0))  # ② spans ②-a and ②-b


def related():
    s, top = slide('related', S1, '관련 연구와 차별점',
                   '같은 최근접 패치 매칭 계열이지만, 관측별 bank와 정상 기준 단위 정렬로 결합을 다르게 한다',
                   note=f'{SRC} §2.1–2.3, 부록 A.4 (서지 검증 2026-10-08).')
    rows = [
        ['Few-shot 패치 매칭 (②-a)', 'PatchCore, AnomalyDINO,\nSuperAD', '계층 특징을 이어 붙인 bank,\nDINOv2 최근접 거리', '관측별 bank 분리 +\n결합 전 정상 기준 단위 정렬'],
        ['다중 척도·다계층 표현 (①)', 'FPN, S², FeatUp,\nDPT, Raghu 2021', '다중 척도·계층 표현으로\ndense prediction 개선', '학습 없이 해상도별 정상 거리 맵 평균\n(S²는 타일 재조립·채널 연결)'],
        ['점수 정규화·결합 (②-a)', 'Kim 2021 (CAM),\nViM 2022', '점수 맵 정규화, 이질적 점수의\n척도를 맞춘 뒤 결합', '정상 support LSO 거리의 q95로\n단위 결정 (레이블 불필요)'],
        ['분포 변화·저-FPR 평가 (문제)', 'ImageNet-C, Guo 2017,\nOvadia 2019, MVTec AD 2', 'shift 아래 성능·신뢰도 평가,\n조명 변화 산업 벤치마크', '저-FPR 곡선 지표와\n고정 임계값 결과를 구분해 평가'],
    ]
    table(s, .70, top + .1, [2.6, 2.8, 3.1, 3.44], ['연구 흐름', '대표 연구', '한 일', 'ScopeAD와의 차이'], rows,
          size=12, bottom=6.4)


# ---------------------------------------------------------------- 02 method
def overview():
    s, top = slide('overview', S2, 'ScopeAD 개요',
                   '정상 support로 48개 bank와 단위 c_m을 한 번 구하고(A), query마다 정렬 → 가중 결합 → 평균한다(B)',
                   note='그림 1 (PaperWorks/figures/fig_scopead_overview.png, 2026-10-12 14:46 수정본). Heatmap은 캐시된 bag 평균 계층별 정규화 맵에서 '
                        f'linear0125로 재구성한 표시용 예시(r = 672 표시), GT 윤곽은 최종 맵 표시용, CQ는 개요도에서 제외. {SRC} §3.1',
                   notes='A Support fitting: 동결 DINOv2, 계층 {5, 11, 23} × 4 해상도 × 4 bag = 48개 bank, LSO 1-NN 정상 거리의 q95로 c_m. '
                         'B Query inference: 4 해상도 × 3 계층 1-NN 거리 → z_m = d_m / c_m. C Support-calibrated layer fusion: '
                         '(r, b)마다 S = z_d + 0.125·z5 (발표 본문의 Deep-Anchored Weighting). D Map aggregation: 4-bag 평균 → bilinear 역투영 '
                         '→ 4-scale 평균, 이미지 점수 = 상위 1% 픽셀 평균.')
    figure(s, FIG / 'fig_scopead_overview.png', top + .02)

def steps():
    s, top = slide('steps', S2, '피팅과 추론의 2단계 절차',
                   '피팅은 support마다 한 번, 추론은 query마다 같은 규칙으로 수행하며 학습되는 값은 없다',
                   ['입력은 정상 support N장, query, 동결 DINOv2-L/14이고, 출력은 픽셀 결함 맵 A(x^q)와 이미지 점수 s(x^q)'],
                   note='피팅: 컴포넌트 m = (l, r, b) 48개 정의 → 전용 bank B_m 생성 → LSO 정상 거리 D_m 측정 → c_m = Q0.95(D_m). '
                        '추론: 4 해상도 특징 추출 → 48개 bank 1-NN 거리 d_m → z_m = d_m / c_m → (r, b)마다 S = z_d + 0.125·z5 → '
                        f'bag 평균 → 해상도 보간 평균 → A, s. 컴포넌트는 같은 support·backbone을 공유하며 독립성을 가정하지 않음. {SRC} §3.1')
    figure(s, CH / 'two_stage.png', top + .1)


def bankcmp():
    s, top = slide('bankcmp', S2, '메모리 뱅크 구축 방식의 차이',
                   '기존 방법은 특징을 통합해 bank 하나를 만들고, ScopeAD는 특징·해상도·bag마다 bank를 따로 만든다',
                   ['기존: 계층 특징을 이어 붙인 bank 1개 → 1-NN 거리 1개라 단위 정렬은 필요 없지만 계층별 정상 변동이 섞임',
                    'ScopeAD: 층 3 × 해상도 4 × bag 4 = 48개 bank를 두고, bank마다 1-NN 거리 맵을 따로 계산',
                    '결합: bank마다 정상 거리 분포가 다르므로 c_m = Q0.95(LSO 거리)로 단위를 맞춘 뒤 z_d + 0.125·z5로 더함'],
                   note='기존 방법은 Outline §1.2·§2.1의 서술(계층 특징 concat, coreset 메모리)과 표 3 기준. 같은 저장 용량에서 단일 bank 57.48% vs '
                        f'관측별 bank 59.57% (표 8 (d), AD2 16-shot seed 0). {SRC} §3.1–3.3')
    figure(s, CH / 'bank_schematic.png', top + .1)


def resol():
    s, top = slide('resol', S2, '관측 해상도와 계층 구성',
                   '패치 크기는 14 px로 고정되므로, 해상도가 높을수록 같은 영역을 더 많은 토큰으로 관찰한다',
                   ['해상도: 짧은 변을 448 / 672 / 896 / 1120 px로 맞춰, 토큰 1개가 덮는 길이가 3.1% → 1.25%로 줄어듦',
                    '계층: DINOv2-L/14의 5·11·23번째 블록 출력 (5 = Shallow, 11·23 = Deep), 모든 query에 같은 구성'],
                   note='아래: 같은 확대 영역(짧은 변의 1/4)을 덮는 토큰 격자. 정사각 기준 토큰 수 1,024 / 2,304 / 4,096 / 6,400. 높은 해상도는 '
                        f'기존 영상의 세밀한 표본화이며 새로운 정상 변동을 만들지 않음. {SRC} §3.2')
    figure(s, CH / 'resolution_grids.png', top + .15)


def bagging():
    s, top = slide('bagging', S2, '메모리 배깅 절차',
                   '이미지당 256토큰을 4번 독립 추출해 bag 4개를 만들고, bag마다 별도 bank와 단위를 둔다',
                   ['① 해상도 r에서 support 이미지마다 패치 토큰 256개를 무작위 비복원 추출',
                    '② 서로 다른 난수 시드 4개로 반복 → 독립 bag b ∈ {1, 2, 3, 4}',
                    '③ 같은 bag 안에서 세 계층은 같은 공간 위치의 토큰을 공유',
                    '④ (l, r, b) 조합 = 컴포넌트 m, 총 48개 · 16-shot bank 1개 = 16 × 256 = 4,096토큰'],
                   note=f'아래 그림 M1: hazelnut query(1120 px)의 bag별 점수 맵, 4-bag 평균, 픽셀별 bag 간 표준편차(평균 0.038). {SRC} §3.2')
    figure(s, CH / 'fig_rsg4_mech_bagging_nocap.png', top + .05)


def whybag():
    s, top = slide('whybag', S2, 'Bag를 사용하는 이유',
                   'bag는 메모리를 줄이는 부분 추출과, 그 추출이 만드는 표본 우연성을 평균으로 상쇄하는 장치다',
                   ['메모리: 모든 토큰 대신 이미지당 256토큰만 저장 → 16-shot bank 1개 = 4,096토큰',
                    '표본 우연성: 한 번만 뽑으면 어떤 정상 토큰이 뽑혔는지에 따라 거리가 흔들림 → 시드가 다른 bag 4개의 점수를 평균',
                    '단위 정렬: bag마다 뽑힌 토큰이 달라 정상 거리 분포도 다름 → bag도 컴포넌트로 보고 c_m을 따로 구함',
                    '효과: 그림 M1에서 bag 간 픽셀 표준편차 평균 0.038, 4-bag 평균이 희소 표본화 차이를 완화'],
                   note='목적 (§3.2): "VRAM 사용량을 낮추고 표본화 분산을 줄이면서 다층·다해상도 패치 특징을 저장하는 참조 뱅크를 구축한다". '
                        '참고: 1120 px에서 모든 토큰을 저장하면 정사각 기준 6,400 × 16장 = 102,400토큰/계층 (§3.2 표로 계산). '
                        f'그림의 선택 위치는 예시용 무작위 추출(실제 실행 인덱스 아님). {SRC} §3.1–3.2, §3.5')
    figure(s, CH / 'bag_sampling.png', top + .05)


def scale_eq():
    s, top = slide('scale_eq', S2, 'Scale Calibration 계산 절차',
                   '각 bank의 정상 LSO 거리 분포에서 q95를 구해, query 거리를 그 단위로 나눈다',
                   note=f'입력: bank B_m, query 특징 f_m(x^q, p), q (기본 0.95). 출력: 단위 척도 c_m, 무차원 점수 맵 z_m. {SRC} §3.1 공통 연산자, §3.3')
    figure(s, CH / 'eq_scale.png', top + .05)


def scale_read():
    s, top = slide('scale_read', S2, 'Scale Calibration 해석',
                   'z = 1은 그 bank의 정상 상위 5% 경계이며, 결합 전 관측 간 비중을 맞추는 기준이다',
                   ['z = 1.0 / 2.0: 패치 거리가 정상 상위 5% 경계와 같음 / 그 2배 — 결합 전 support 참조 단위',
                    'support 기준점: LSO 정상 거리 표본에서 z ≤ 1인 비율 ≈ 95%, 최종 판정 임계값은 결합 맵에서 따로 정함',
                    '예시 (오른쪽): 같은 d = 0.178이 Layer 5 기준(0.123)은 넘고 Layer 23 기준(0.232)은 못 넘음'],
                   note=f'그림 M2: hazelnut 1120 px, bag 평균, d*는 예시값. 파라미터 상태: q 민감도 표 11, 정규화 방식 비교 표 8c·8d. {SRC} §3.3')
    figure(s, CH / 'fig_rsg4_mech_scale_nocap.png', top + .02)


def control_eq():
    s, top = slide('control_eq', S2, 'Deep-Anchored Weighting 계산 절차',
                   'Deep 평균을 기준선으로 두고, 정규화된 Shallow 점수를 고정 가중치 0.125로 더한다',
                   ['규격: 임계값·게이트·학습되는 가중치 없음, 계층마다 bank와 점수 맵을 유지한 채 점수 단위에서 결합',
                    'λ 선택: AD2 2-shot seed 0 AUPRO로 {0, 0.125, 0.25, 0.5, 1, 2} 중 선택해 동결 (post-hoc, 재조정 없음)',
                    '주의: z5는 정규화 뒤 점수라 λ의 상대 기여는 정규화 방식에 따라 다름 → 방식 간 같은 λ를 직접 비교하지 않음'],
                   note=f'목적: Deep 기준선에 shallow 관측을 고정 가중치로 가산해 관측 축을 깊이 방향으로 확장. 입력: z_l,r,b (l = 5, 11, 23), λ = 0.125. {SRC} §3.4')
    figure(s, CH / 'eq_control.png', top + .15, bottom=min(6.6, top + 2.6))


def control_ex():
    s, top = slide('control_ex', S2, 'Deep-Anchored Weighting 예시',
                   'Shallow 가산분은 0.05–0.375로 작아, 정상 면의 점수를 크게 올리지 않는다',
                   ['예시 값: 정상 질감 변화 면은 0.806, 미세 스크래치는 1.200 — z5 전체를 더하면 각각 1.90과 3.65까지 오름',
                    '실제 맵: 0.125·z5는 같은 색 범위에서 거의 보이지 않을 만큼 작고, λ = 1이면 정상 껍질까지 밝아짐'],
                   note=f'위: Outline §3.4 예시 계산(실측 아님, 전체 가산 = z_d + z5). 아래: hazelnut 1120 px, bag 평균, 색 범위 [0, 2]. {MAPNOTE}. {SRC} §3.4')
    stacked(s, [CH / 'fusion_example.png', CH / 'fusion_maps.png'], top + .05)


def aggregate():
    s, top = slide('aggregate', S2, '맵 통합과 이미지 점수',
                   'bag 평균 → 4개 해상도 보간 평균으로 픽셀 맵을 만들고, 상위 1% 픽셀 평균을 이미지 점수로 쓴다',
                   note=f'입력: 각 (r, b)의 결합 점수 맵 S_r,b, R = {{448, 672, 896, 1120}}. 출력: A(x^q) ∈ R^(H×W), s(x^q). {SRC} §3.5')
    figure(s, CH / 'eq_aggregate.png', top + .1)


def spec():
    s, top = slide('spec', S2, '파라미터 규격과 재현성',
                   '모든 값은 사전에 고정하고, 추론에는 역전파·파라미터 갱신·테스트 라벨을 쓰지 않는다',
                   ['백본은 동결 DINOv2-L/14(약 304M, 패치 14)이고, support 수 N은 1, 2, 4, 8, 16장이다',
                    '관측은 계층 {5, 11, 23}, 짧은 변 {448, 672, 896, 1120}, bag 4개 × 이미지당 256토큰으로 고정한다',
                    '결합은 q = 0.95와 λ = 0.125(Layer 5에만)이며, λ는 AD2 2-shot seed 0로 고른 뒤 고정했다',
                    '계산은 FP16 autocast, bank 256토큰/source이며, 선택적으로 연결 성분 최댓값 채움(CQ) 후처리를 쓴다',
                    '모든 실험은 seed 0·1·2와 고정 support manifest로 결정론적으로 실행된다'],
                   note=f'표 8 등 ablation 실행 조건은 CQ99. {SRC} §3.6')
    boxed(s, .70, top + .35, 11.94, .62, '학습되는 값 없음',
          '임계값·게이트·학습 가중치 없이, 정상 support에서 계산한 bank와 단위 c_m만 query 추론에 쓴다')


def t1():
    s, top = slide('t1', S3, '평가 데이터 (표 1)',
                   '주 평가는 여러 조명이 섞인 MVTec AD 2 test_public이고, support는 regular 조명에서 고른다',
                   ['Shift 1–3 조명은 support에 없던 변화 → 조명 변화 아래의 국소화를 함께 평가',
                    'test_private,mix는 대응 장면의 seen/unseen 조명을 섞은 별도 split (public·private 합집합 아님)'],
                   note=f'{SRC} §4.1 표 1')
    raw_table(s, '**표 1.', top + .1, [2.6, 3.4, 5.94], header=['Split', '이미지 수', '조명 · 레이블 범위'], size=12,
              bottom=6.0)


def t2():
    s, top = slide('t2', S3, '평가 설정 (표 2)',
                   '주 지표는 AUPRO@0.05이며, 연속 점수 지표는 임계값 없이 점수 맵에서 바로 계산한다',
                   note=f'{SRC} §4.1 표 2')
    raw_table(s, '**표 2.', top + .1, [2.6, 9.34], header=['항목', '설정'], size=11, bottom=6.4)


def t3():
    s, top = slide('t3', S3, '비교군 설정 (표 3)',
                   '비교군은 유효 입력을 1120으로 바꿔 다시 실행했고, 백본·계층·후처리는 각 방법 설정을 유지한다',
                   note='모든 방법은 같은 support 경로·query·evaluator 사용. FastRecon·RAD는 논문 식의 독립 구현, SuperAD는 공통 support 사용 '
                        f'(원 논문은 전체 train coreset) → 외부 비교는 전체 파이프라인 비교. {SRC} §4.1 표 3')
    raw_table(s, '**표 3.', top + .05, [2.3, 3.0, 2.8, 3.84],
              header=['방법', '학습 · support fit', 'Backbone / feature layers', '입력 · resolution · 추가 설정'], size=10,
              bottom=6.72)


def t4():
    s, top = slide('t4', S3, 'Few-shot 주 비교 (표 4)',
                   '1–16 shot 세 지표 모두 1위, 16-shot 3-seed AUPRO@0.05는 SuperAD보다 +4.17pp',
                   note='AUPRO@0.05는 비교군 최고값 대비 +3.24–8.50pp (1-shot 최대). 비교군: seed 0–2 8-class macro의 mean ± sample SD, 1120 입력 FP32. ‡ ScopeAD: 2·16-shot 3-seed, 1·4·8-shot seed 0 단일값, '
                        f'FP16·4 해상도. † RegAD는 공식 STN 구조상 224. 굵게 = 최고, 밑줄 = 차상위. {SRC} §4.2 표 4')
    raw_table(s, '**표 4.', top + .02, [.6, 1.3] + [1.255] * 8,
              header=['Shot', 'Metric', 'PatchCore', 'AnomalyDINO', 'FastRecon', 'RAD', 'SubspaceAD', 'SuperAD',
                      'RegAD (224)†', 'ScopeAD‡'], size=10, row_h=.26, strong_cols=(0,), tight=True)


def t5():
    s, top = slide('t5', S3, 'MVTec AD · VisA 비교 (표 5)',
                   'I-AUROC는 10개 조건 중 7개에서 1위, P-AUROC는 최고값과 0.04–1.36pp 이내다',
                   ['I-AUROC 2위 조건: MVTec AD 16-shot, VisA 1·2-shot (모두 SuperAD 다음)'],
                   note='seed 0 단일, support hash 일치 클래스 macro (MVTec AD 15, VisA 12). 각 셀 I-AUROC / P-AUROC(전체 FPR). '
                        f'FastRecon = E9 DINOv2 기반 변형. linear0125 hash 재감사는 수집 단계 과제. {SRC} §4.2 표 5')
    raw_table(s, '**표 5.', top + .05, [1.0, .62] + [1.29] * 8,
              header=['Dataset', 'Shot', 'PatchCore', 'AnomalyDINO', 'FastRecon', 'RAD', 'SubspaceAD', 'SuperAD', 'RegAD',
                      'ScopeAD'], size=10, row_h=.36)


def t5b():
    s, top = slide('t5b', S3, '3-seed 결과 (표 5b)',
                   '3-seed 평균도 seed 0와 비슷한 수준이다 (16-shot I-AUROC MVTec AD 98.52, VisA 95.78)',
                   ['seed 0 대비: MVTec AD 16-shot I-AUROC 98.74 → 3-seed 98.52, VisA 95.99 → 95.78',
                    '주의: seed 1·2의 support는 비교군과 짝지어지지 않아 표 5의 비교군과 직접 순위를 매기지 않음'],
                   note=f'mean ± sample SD, 각 셀 I-AUROC / P-AUROC (%). {SRC} §4.2 표 5b')
    raw_table(s, '**표 5b.', top + .3, [2.4, 3.9, 3.9], x=1.6, header=['Dataset', '2-shot', '16-shot'], size=14,
              bottom=top + 1.9)


def t6():
    s, top = slide('t6', S3, '공식 private 평가 (표 6)',
                   'Ours 행은 이전 게이트 설정의 제출 결과이며, linear0125로 같은 절차로 재제출해 교체해야 한다',
                   note='각 셀 TESTpriv / TESTpriv,mix. † Layer 5 β = 0.99 초과 잔차를 더한 이전 설정. ClassF1·SegF1 임계값: support 제외 '
                        'train/good 8번째 간격 샘플 pixel p95 × 1.421 (private label 미사용). Full-data 행은 원 논문·리더보드 수치로 '
                        f'데이터 예산·백본·입력·후처리가 달라 순위화하지 않음. {SRC} §4.3.1 표 6')
    raw_table(s, '**표 6.', top + .02, [4.44, 2.5, 2.5, 2.5], header=['방법', 'AU-PRO@0.05', 'ClassF1', 'SegF1'], size=10,
              row_h=.24, tight=True)


def t7():
    s, top = slide('t7', S3, '분포 변화 평가 (그림 2 · 표 7)',
                   '조명 strata(그림 2)와 합성 손상(표 7)은 linear0125로 아직 평가하지 않았다',
                   ['그림 2: 기존 그림의 ScopeAD 선은 게이트 결과라 사용하지 않음 → condition strata 집계를 켜고 재평가',
                    '표 7: 비교군만 확보, Gaussian noise에서 PatchCore 74.03 / 83.66이 가장 높음'],
                   note='각 셀 AUPRO@0.05 / pAUROC@0.05. MVTec AD 15클래스, 16-shot, seed 0, severity 3, support는 깨끗한 이미지. '
                        f'† DINOv2 기반 프로젝트 변형(표 4의 fastrecon_paper와 다른 구성). {SRC} §4.3.2–4.3.3 표 7')
    raw_table(s, '**표 7.', top + .05, [2.94, 1.8, 1.8, 1.8, 1.8, 1.8],
              header=['방법', 'Clean', 'Brightness s3', 'Contrast s3', 'Defocus blur s3', 'Gaussian noise s3'], size=10,
              row_h=.36, mod=lambda r: [r[0].replace(' (project adaptation)', ''), *r[1:]])


# ---------------------------------------------------------------- 04 ablation (outline tables as-is)
def t8():
    s, top = slide('t8', S4, '구성 요소 소거 (표 8)',
                   '어느 구성을 하나만 바꿔도 Full(59.57)보다 낮고, 해상도와 bank 구조의 영향이 가장 크다',
                   ['관측: 448 단일 해상도 −7.56pp / 결합: 단일 bank −2.09, 3계층 동일 가중 −1.36, MAD −1.05pp',
                    '단일 bank는 같은 저장 float 수·코사인 거리 정의를 유지 → bank 구조 차이만 분리 (pAUROC −2.66pp)',
                    '정규화 생략은 AUPRO −0.40pp, pAUROC +0.19pp → q95의 정규화 생략 대비 이득은 주장하지 않음'],
                   note='동결 DINOv2 ViT-L/14, FP16, 4 해상도, cap256, 4 bags, CQ99, top-1% mean 이미지 점수, 동일 support. 각 행은 Full에서 '
                        f'한 구성만 바꿈. 단일 seed 개발 비교. {SRC} §4.4.1 표 8')
    raw_table(s, '**표 8.', top + .05, [3.0, 1.6, 1.6, 1.6, 1.3, 1.3, 1.54],
              header=['구성', 'Bank 구조', 'Score normalization', 'Shallow 항', 'AUPRO@0.05', 'pAUROC@0.05', 'Δ AUPRO vs Full'],
              size=11, bottom=6.6)


def t8a():
    s, top = slide('t8a', S4, '데이터셋 · shot별 구성 소거 (표 8a)',
                   '세 데이터셋 모든 shot에서 Full이 3계층 동일 가중 평균과 계층 concat보다 높다',
                   ['정규화 생략: MVTec AD·VisA I-AUROC는 0.07–1.06pp 높고 AD2 AUPRO는 −0.40 ~ +0.33pp로 엇갈림'],
                   note='AD2 = AUPRO@0.05 / pAUROC@0.05, MVTec AD·VisA = I-AUROC / P-AUROC(전체 FPR). seed 0 단일값. AD2 1/4/8-shot은 로컬 4090, '
                        f'2/16-shot Full은 fusion sensitivity 실행의 동일 설정. {SRC} §4.4.1 표 8a')
    raw_table(s, '**표 8a.', top + .02, [1.7, .9, 2.35, 2.35, 2.35, 2.29],
              header=['Dataset', 'Shot', 'Full (linear0125)', 'No normalization', 'Equal-weight mean', 'Layer concat'],
              size=10, row_h=.27)


def t8b():
    s, top = slide('t8b', S4, 'Shallow 가중치 민감도 (표 8b)',
                   'λ ≥ 0.5에서 AUPRO@0.05가 빠르게 낮아지고, λ = 0.125가 2·16-shot 모두 가장 높다',
                   ['16-shot: λ = 0.125 59.57 → λ = 1 54.02 (−5.55pp) → λ = 2 47.33',
                    '선택 절차: 2-shot AUPRO만으로 λ를 고른 뒤 16-shot 해석 전에 동결 (개발 평가 사용, post-hoc)'],
                   note=f'd = (z11 + z23)/2, s = z5. FP16, 4-scale, cap256, q95, CQ99 고정, seed 0. 표에는 양의 가중치 조건만 제시. {SRC} §4.4.2 표 8b')
    raw_table(s, '**표 8b.', top + .15, [3.94, 2.0, 2.0, 2.0, 2.0],
              header=['결합 규칙', '2-shot AUPRO', '2-shot pAUROC', '16-shot AUPRO', '16-shot pAUROC'], size=12, bottom=5.6)


def t8c():
    s, top = slide('t8c', S4, '정규화 방식 비교 MVTec AD · VisA (표 8c)',
                   'q95는 MAD보다 모든 shot에서 I-AUROC가 높지만, 정규화 생략(none)보다는 0.07–1.06pp 낮다',
                   ['q95 − MAD: MVTec AD I-AUROC +0.47 ~ +1.60pp, VisA +0.34 ~ +1.30pp',
                    'P-AUROC: MVTec AD에서만 q95가 none보다 +0.09 ~ +0.18pp'],
                   note=f'전체 방법(linear0125)에서 normalization만 바꾼 대조, 4 해상도, cap256, 같은 support, seed 0. 각 셀 I-AUROC / P-AUROC. {SRC} §4.4.1 표 8c')
    raw_table(s, '**표 8c.', top + .05, [2.3, 1.2, 2.8, 2.8, 2.84], header=['Dataset', 'Shot', 'none', 'MAD', 'q95 (기본)'],
              size=11, row_h=.32)


def t8d():
    s, top = slide('t8d', S4, '정규화 방식 비교 AD2 (표 8d)',
                   'AD2에서 q95는 MAD보다 AUPRO@0.05 +0.76–1.69pp 높고, none과는 −0.33 ~ +0.40pp로 엇갈린다',
                   ['q95 − none: AUPRO 평균 +0.02pp, pAUROC는 4개 shot에서 q95가 0.19–0.52pp 낮음',
                    'LOO-token vs LSO: AD2 AUPRO 59.45 vs 59.33 (2-shot), 59.55 vs 59.57 (16-shot) → LSO 우위 주장 안 함'],
                   note='각 셀 AUPRO@0.05 / pAUROC@0.05 / I-AUROC. 4 해상도, cap256, 같은 support. none·q95 열은 dsba3 재실행으로 로컬 ablation과 '
                        f'소수 둘째 자리까지 일치 확인. {SRC} §4.4.1 표 8d, calibration support 제외 단위 대조')
    raw_table(s, '**표 8d.', top + .15, [.9, 2.75, 2.75, 2.75, 1.4, 1.39],
              header=['Shot', 'none', 'MAD', 'q95 (기본)', 'q95 − none (AUPRO)', 'q95 − MAD (AUPRO)'], size=12, bottom=5.7)


def t10():
    s, top = slide('t10', S4, '총 bank 용량 통제 해상도 비교 (표 10)',
                   '총 토큰 수를 맞춰도 다해상도가 1120 단일보다 +2.52pp 높다 — 이전 게이트 설정의 결과',
                   ['용량 통제: 1120 단일에 토큰을 4배(1024) 줘도 56.63%로 오히려 −0.33pp',
                    '해석: 최고 해상도와 총 bank 용량만으로 다해상도 결합의 이득을 설명할 수 없음',
                    '주의: 이전 비대칭 게이트 설정(E12, RSG3) 결과 → linear0125로 같은 비교를 다시 실행해 교체 예정'],
                   note=f'AD2 test_public, 16-shot, seed 0, full8, DINOv2 ViT-L/14. {SRC} §4.4.3 표 10')
    raw_table(s, '**표 10.', top + .1, [3.14, 1.3, 1.3, 1.2, 1.3, 1.85, 1.85],
              header=['해상도 구성', '토큰 수 / Bag', 'Full8 AUPRO@0.05', 'Δ vs 단일', 'Full8 pAUROC@0.05',
                      'Regular AUPRO / pAUROC', 'Shift 1 AUPRO / pAUROC'], size=11, bottom=5.6)


def t10a():
    s, top = slide('t10a', S4, '단일 · 다중 해상도 비교 (표 10a)',
                   '4개 해상도 결합은 최선의 단일 해상도보다 AUPRO@0.05가 +4.95pp(2-shot), +2.10pp(16-shot) 높다',
                   ['단일 해상도는 높을수록 좋지 않음: 2-shot 최선 672, 16-shot 최선 1120 (1120 대비 2-shot +5.06pp)',
                    '주의: bank당 cap은 같지만 총 용량은 달라, 총 용량 효과는 이 표로 분리하지 않음 (→ 표 10)'],
                   note=f'linear0125, q95, FP16, cap256, 같은 support, seed 0, full8. 각 셀 AUPRO@0.05 / pAUROC@0.05. {SRC} §4.4.3 표 10a')
    raw_table(s, '**표 10a.', top + .15, [4.4, 3.0, 3.0], x=1.3, header=['해상도 구성', '2-shot', '16-shot'], size=13,
              bottom=5.9)


def t10b():
    s, top = slide('t10b', S4, 'MVTec AD · VisA 해상도 비교 (표 10b)',
                   '전체 방법에서 4 해상도 결합이 두 데이터셋 · 두 shot의 I-AUROC와 P-AUROC 모두 1120 단일보다 높다',
                   ['I-AUROC 이득: MVTec AD +1.23 / +0.53pp, VisA +5.27 / +1.60pp (2 / 16-shot)',
                    '주의: 총 bank 용량 효과는 분리하지 않음'],
                   note=f'linear0125, q95, FP16, cap256, seed 0. 각 셀 I-AUROC / P-AUROC. {SRC} §4.4.3 표 10b')
    raw_table(s, '**표 10b.', top + .2, [2.4, 2.6, 2.7, 2.7], x=1.3, header=['Dataset', '해상도 구성', '2-shot', '16-shot'],
              size=13, bottom=5.4)


def t11():
    s, top = slide('t11', S4, '척도 분위수 민감도 (표 11)',
                   'q = 0.80–0.999에서 AUPRO@0.05 변화는 −0.22 ~ +0.04pp로 평탄하다',
                   ['중앙값 기준(q = 0.5)만 낮고, 꼬리확률 변환·토큰별 로컬 척도도 두 shot에서 일관된 개선 없음',
                    '척도 진단: 12개 컴포넌트 q95 척도 차이 평균 약 1.8배(deep 사이 약 1.4배), 변동계수 0.18–0.24'],
                   note='잔차 게이트 분위수 β는 방법에서 제거. λ = 0.125 고정, seed 0, AD2. 후속: local-center 보정(AD2 2·16-shot, 강도 0.25/0.5) '
                        f'검증 중, 확정 전에는 성능 기여로 기술하지 않음. {SRC} §4.4.4 표 11·척도 진단')
    raw_table(s, '**표 11.', top + .05, [4.6, 2.3, 2.3, 2.74],
              header=['척도 분위수 q', '2-shot', '16-shot', 'Δ AUPRO vs Default (2 / 16-shot)'], size=11, row_h=.33)


def t12():
    s, top = slide('t12', S4, '효율 · 정성 분석 (표 12 · 12a · 그림 3)',
                   'linear0125의 계산 비용과 정성 비교 그림은 아직 계측·작성하지 않았다',
                   ['표 12·12a: 게이트 설정 측정값은 제외, bank 벡터 수만 구성에서 계산 (16-shot, 컴포넌트당 4,096토큰)',
                    '그림 3 사양: Regular·Shift 1·Gaussian noise 3행 × (Input, GT, SuperAD, PatchCore, ScopeAD, λ·z5)'],
                   note='SuperAD 측정 조건: AD2 full8 test_public, seed 0, 16-shot, DINOv2 ViT-L/14, FP32, batch 1. 구성 요소별 시간 분해도 '
                        f'linear0125에서 재계측 필요. {SRC} §4.5 표 12·12a, 그림 3')
    raw_table(s, '**표 12.', top + .1, [1.9, 2.2, 1.4, 1.5, 1.6, 1.7, 1.64],
              header=['Arm', 'Bank vectors / class', 'Fit (s)', 'Query mean (ms)', 'Pooled mean / p95 (ms)',
                      'Peak alloc fit / query (GiB)', 'Peak GPU fit / query (GiB)'], size=10, row_h=.4)
    y = table_bottom(top + .1, 3, .4) + .3
    raw_table(s, '**표 12a.', y, [1.4, 2.4, 2.6, 2.8, 2.74],
              header=['Cap', 'Fit (s/class)', 'Query mean (s/image)', '클래스별 query p95 (s/image)', 'Peak alloc fit / query (GiB)'],
              size=10, row_h=.34, name='SummaryTable2')


# ---------------------------------------------------------------- 05 discussion
def discussion():
    s, top = slide('discussion', S5, '논의',
                   '다해상도와 관측별 bank는 서로 다른 단계에서 기여하고, Shallow 가산은 가중치에 민감하다',
                   [f'관측 확장: 최선 단일 해상도 대비 +4.95 / +2.10pp, 용량 통제(게이트 설정)에서도 +2.52pp ({P["t10a"]}·{P["t10"]}쪽)',
                    f'관측별 결합: 세 데이터셋 모든 shot에서 단일 bank보다 높음 ({P["t8a"]}쪽)',
                    f'q95 정렬: MAD보다 높지만 정규화 생략과 차이 작음 → 척도 차이가 큰 관측·조건 변화 query에서 확인 ({P["t8d"]}쪽)',
                    '배포: 단위 정렬이 query FPR을 정하지 않음 → 조명 조건별 임계값 보정과 검증이 필요'],
                   note=f'{SRC} §5.1–5.2, 부록 A.1.')
    boxed(s, .70, top + .3, 11.94, 1.2, '주장 범위',
          '지지(개발 모집단): 관측별 bank > 단일 bank, q95 > MAD, Shallow 비중 조절 필요, 다해상도 이득, AD2 2·16-shot SuperAD 대비 우위. '
          '미검증: q95의 정규화 생략 대비 이득, linear0125의 private·shift·corruption·계산 비용.')


def pending():
    s, top = slide('pending', S5, '남은 검증',
                   '게이트 설정으로 남은 표 6·10을 교체하고, 분포 변화·비용·다중 seed를 채운다',
                   note=f'{SRC} 부록 A.1–A.3, §4.4.4 후속 실험 상태, §4.5.')
    rows = [
        ['공식 private (표 6)', '게이트 설정 결과 64.80 / 62.49', 'linear0125, 같은 support·임계값 절차로 재제출'],
        ['용량 통제 해상도 비교 (표 10)', '게이트 설정 결과', 'multi-scale vs 1120 (cap256 / 1024) 재실행'],
        ['조명 strata (그림 2)', 'ScopeAD 선 미사용', 'condition strata 집계를 켜고 재평가'],
        ['합성 손상 (표 7)', 'ScopeAD 미실행', 'MVTec AD 15클래스 × 5조건 실행'],
        ['계산 비용 (표 12 · 12a)', 'ScopeAD 미계측', 'linear0125 파이프라인 재계측 (구성 요소별 시간 포함)'],
        ['다중 seed', '2·16-shot 3-seed 완료', '1·4·8-shot seed 1·2'],
        ['local-center 보정', '진행 중 (AD2 2·16-shot)', 'q95·none을 모두 넘는지 확인 후 판단'],
        ['원고 그림 M3 · 3', 'M3는 게이트 맵, 그림 3은 사양만 확정', 'linear0125 기준 Deep·Shallow·결합 맵으로 작성'],
    ]
    table(s, .70, top + .1, [3.4, 3.6, 4.94], ['항목', '현재 상태', '다음 단계'], rows, size=12, bottom=6.5)


def conclusion():
    s, top = slide('conclusion', S5, '결론 및 향후 진행 방향',
                   '관측을 늘리고 관측별 점수를 맞춰 결합하는 ScopeAD로 AD2 few-shot 국소화 1위를 얻었다',
                   note=f'{SRC} §6, 부록 A.1. 근거 쪽: 관측 확장 {P["t10a"]}, 결합 {P["t8"]}·{P["t8d"]}, 성능 {P["t4"]}, 남은 항목 {P["pending"]}.')
    rows = [
        ['관측 확장', '4개 해상도 + Shallow / Deep 관측,\n최선 단일 대비 +4.95 / +2.10pp', '용량 통제 비교는 게이트 결과', '표 10 linear0125 재실행'],
        ['관측별 결합', '관측별 bank + q95 정렬 + λ = 0.125,\n단일 bank 대비 +2.09pp', 'q95의 정규화 생략 대비 이득 없음', 'local-center 등 정렬 방식 검증'],
        ['성능', 'AD2 16-shot 3-seed 60.06 ± 0.88,\nSuperAD 대비 +4.17pp', '1·4·8-shot은 seed 0 단일값', '다중 seed 완성'],
        ['강건성 · 배포', '— (private은 게이트 설정 결과)', 'linear0125 private·shift·corruption·비용 미확보', '재제출 · 재평가 · 재계측'],
    ]
    table(s, .70, top + .1, [1.8, 4.0, 3.2, 2.94], ['구분', 'Summary', '현재 문제', '향후 진행 방향'], rows, size=13,
          bottom=6.2)


def a_claims():
    s, top = slide('a_claims', SA, 'Claim–evidence–status matrix',
                   '원고의 각 주장과 근거, 현재 지지 상태를 정리한 저자용 표 (Outline 부록 A.1 그대로)',
                   note=f'{SRC} 부록 A.1')
    raw_table(s, '| Claim', top + .02, [2.6, 5.4, 1.5, 2.44], header=['Claim', '핵심 근거', '상태', '원고 내 기술 원칙'],
              size=10, bottom=6.72, plain=False)


for key in ORDER:
    globals()[key]()

groups = [('표지 · 요약', ['cover', 'toc', 'summary']),
          ('01 Introduction', ['bg', 'ambig', 'align', 'p1', 'p2a', 'p2b', 'contrib', 'related']),
          ('02 Method', ['overview', 'steps', 'bankcmp', 'resol', 'bagging', 'whybag', 'scale_eq', 'scale_read',
                         'control_eq', 'control_ex', 'aggregate', 'spec']),
          ('03 실험 결과', ['t1', 't2', 't3', 't4', 't5', 't5b', 't6', 't7']),
          ('04 Ablation', ['t8', 't8a', 't8b', 't8c', 't8d', 't10', 't10a', 't10b', 't11', 't12']),
          ('05 논의', ['discussion', 'pending', 'conclusion']),
          ('Appendix A', ORDER[MAIN_LAST:])]
assert [k for _, ks in groups for k in ks] == ORDER
set_sections(prs, [(n, [slides[k] for k in ks]) for n, ks in groups])
prs.save(OUT)
print('saved', OUT, len(prs.slides), 'slides')
