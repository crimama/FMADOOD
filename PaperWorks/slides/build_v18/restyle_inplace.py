"""Apply the plain black-line table style to result tables (section 03/04 slides) of an existing deck in place.

Only table XML of those slides is touched; every other shape, slide, note and the user's edits stay as they are.
usage: python3 restyle_inplace.py in.pptx out.pptx
"""
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

PLAIN_STYLE = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'  # No Style, No Grid
LINE_W = 9525  # 0.75 pt
RESULT_SECTIONS = ('03.', '04.')


def plain_table(frame):
    tbl = frame._element.graphic.graphicData.tbl
    tblPr = tbl.tblPr
    tblPr.set('firstRow', '0')
    tblPr.set('bandRow', '0')
    sid = tblPr.find(qn('a:tableStyleId'))
    if sid is None:
        sid = etree.SubElement(tblPr, qn('a:tableStyleId'))
    sid.text = PLAIN_STYLE
    drop = {qn(t) for t in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB', 'a:solidFill', 'a:noFill', 'a:gradFill', 'a:blipFill',
                            'a:pattFill', 'a:grpFill')}
    for tc in tbl.iter(qn('a:tc')):
        tcPr = tc.get_or_add_tcPr()
        for child in list(tcPr):
            if child.tag in drop:
                tcPr.remove(child)
        for pos, tag in enumerate(('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB')):
            ln = etree.Element(qn(tag), w=str(LINE_W), cap='flat', cmpd='sng', algn='ctr')
            fill = etree.SubElement(ln, qn('a:solidFill'))
            etree.SubElement(fill, qn('a:srgbClr'), val='000000')
            etree.SubElement(ln, qn('a:prstDash'), val='solid')
            tcPr.insert(pos, ln)
        tcPr.insert(4, etree.Element(qn('a:noFill')))
    for row in frame.table.rows:
        for cell in row.cells:
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.color.rgb = RGBColor.from_string('000000')


def section_of(slide):
    for sh in slide.shapes:
        if sh.name == 'SectionLabel' and sh.has_text_frame:
            return sh.text_frame.text.strip()
    return ''


def main(src, dst):
    prs = Presentation(src)
    done = []
    for i, s in enumerate(prs.slides, 1):
        if not section_of(s).startswith(RESULT_SECTIONS):
            continue
        for sh in s.shapes:
            if getattr(sh, 'has_table', False) and sh.has_table:
                plain_table(sh)
                done.append(i)
    prs.save(dst)
    print('restyled tables on slides:', sorted(set(done)))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
