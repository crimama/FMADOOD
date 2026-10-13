"""Build the ScopeAD (Outline v18) deck in the DSBA template language, structured like the v16.3 deck.

Main deck = Introduction (logic, claims, contributions) -> Method -> results organised by the outline's tables and
figures (charts) -> discussion; exact numbers in Appendix tables parsed from Outline_v18.md.
usage: python3 build_deck_v18.py out.pptx      (charts: run charts_v18.py / draw_overview_scopead.py first)
"""
import os
import re
import sys
from pathlib import Path

sys.path.insert(0, os.path.expanduser('~/.claude/skills/dsba-slide-deck/scripts'))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from deck_style import (BOLD, FONT, INK, LOGO, NAVY, SUB, XBOLD, _textbox, figure, frame_slide,  # noqa: E402
                        new_deck, style_run, table, text_box)
from deck_ops import set_sections  # noqa: E402
from outline_tables import table as otable  # noqa: E402

CH = Path(__file__).resolve().parent / 'charts'
OUT = sys.argv[1]
SRC = '출처: PaperWorks/outlines/Outline_v18.md'
MAPNOTE = ('맵은 PaperWorks/figures/cache_rsg4_mechanism의 계층별 맵을 z_d + 0.125·z5로 다시 합성한 설명용 그림 '
           '(캐시 실행, FP32, 실험 결과 아님)')

ORDER = ['cover', 'toc', 'summary',
         'bg', 'ambig', 'expand', 'hard1', 'hard2', 'split', 'roles', 'contrib', 'related',
         'overview', 'bank', 'scale', 'control', 'maps', 'spec',
         'setup', 'baselines', 't4', 't4aux', 't5', 't5b', 't6', 'shift',
         't8', 't8a', 't8b', 't8cd', 't10ab', 't10', 't11',
         'discussion', 'pending', 'conclusion',
         'a_t4', 'a_t4aux', 'a_t5', 'a_t6', 'a_t7', 'a_t8', 'a_t8a', 'a_t8cd', 'a_t10', 'a_t11']
MAIN_LAST = ORDER.index('conclusion') + 1
P = {k: i + 1 for i, k in enumerate(ORDER)}
A = {k: f'A-{i + 1}' for i, k in enumerate(ORDER[MAIN_LAST:])}


def page(key):
    return A.get(key, P[key])


S1, S2, S3, S4, S5 = ('01. Introduction', '02. Method', '03. 실험 결과', '04. Ablation · 분석', '05. 논의와 향후 계획')
SA = 'Appendix A. 결과 표 원본'
prs = new_deck()
slides = {}


def em(text):
    return sum(1.0 if ord(c) > 0x1100 else (0.3 if c == ' ' else 0.6) for c in text)


def lint(key, title, lead, detail):
    if em(title) * 28 / 72 > 10.6:
        print(f'WARN title wraps  p{page(key)}: {title}')
    if em(lead) * 16 / 72 > 11.4:
        print(f'WARN lead wraps   p{page(key)}: {lead} ({em(lead) * 16 / 72:.1f})')
    for line in detail:
        if em(line) * 16 / 72 > 11.6:
            print(f'WARN bullet wraps p{page(key)}: {line} ({em(line) * 16 / 72:.1f})')


def slide(key, section, title, lead, detail=(), note='', notes=''):
    lint(key, title, lead, detail)
    s, top = frame_slide(prs, section=section, title=title, lead=lead, detail=detail, note=note, notes=notes,
                         page=page(key))
    slides[key] = s
    return s, top


def boxed(s, x, y, w, h, label, text):
    sh = _textbox(s, x, y, w, h, 'KeyBox')
    sh.line.color.rgb = RGBColor.from_string(NAVY)
    sh.line.width = Pt(1)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(.2)
    tf.margin_top = tf.margin_bottom = Inches(.12)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = label + '  '; style_run(r, 15, NAVY, BOLD, True)
    r = p.add_run(); r.text = text; style_run(r, 15, INK, FONT, False)


def pretty(cell):
    c = cell.replace('$', '').replace('lambda', 'λ').replace('Delta', 'Δ').replace(',z_5', '·z5')
    c = c.replace('z_5', 'z5').replace('^*', '*').replace('−', '−')
    return re.sub(r'\s+', ' ', c).strip()


# ---------------------------------------------------------------- cover / toc / summary
def cover():
    blank = next(l for l in prs.slide_layouts if l.name == '빈 화면')
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(str(LOGO), Inches(.9), Inches(1.15), Inches(1.55), Inches(.81)).name = 'Logo'
    text_box(s, .9, 2.25, 11.5, .8, 'ScopeAD', 44, NAVY, XBOLD, True, name='Title')
    text_box(s, .9, 3.10, 11.8, .9, ['Observation Expansion and Support-Calibrated Fusion', 'for Few-Shot Industrial '
             'Anomaly Localization'], 20, NAVY, BOLD, True, name='Subtitle')
    text_box(s, .9, 4.0, 11.5, .5, '소수 정상 이미지에서 관측을 늘리고, 관측별 점수를 정상 기준으로 맞춰 결합하는 결함 국소화', 18,
             INK, FONT, name='Subtitle2')
    line = s.shapes.add_connector(1, Inches(.9), Inches(4.75), Inches(12.4), Inches(4.75))
    line.line.color.rgb = RGBColor.from_string(NAVY)
    line.line.width = Inches(.015)
    text_box(s, .9, 4.95, 11.5, .4, '연구 아웃라인 v18 발표  ·  2026-10-12', 16, SUB, FONT, name='Date')
    s.notes_slide.notes_text_frame.text = ('근거 문서: PaperWorks/outlines/Outline_v18.md. ScopeAD = Support-Calibrated '
                                           'Observation-expanded Patch Ensemble for Anomaly Detection (이전 이름 SCARF).')
    slides['cover'] = s


def toc():
    rng = lambda a, b: f'{P[a]}–{P[b]}쪽'  # noqa: E731
    slide('toc', '목차', '발표 순서', '흐름 │ 문제와 주장(Introduction) → 방법 → 주 결과 → 구성 요소 분석 → 남은 검증',
          [f'01 Introduction: 왜 관측을 늘리고, 왜 관측별로 맞춰 결합해야 하는가 · 기여 ({rng("bg", "related")})',
           f'02 Method: 관측별 bank, Scale Calibration, Contribution Control, 맵 통합 ({rng("overview", "spec")})',
           f'03 실험 결과: 표 4–6, 그림 2·표 7 (few-shot 주 비교, MVTec AD·VisA, private) ({rng("setup", "shift")})',
           f'04 Ablation · 분석: 표 8–8d, 표 10–10b, 표 11 ({rng("t8", "t11")})',
           f'05 논의와 향후 계획: 주장 범위와 남은 검증 ({rng("discussion", "conclusion")})',
           f'Appendix A: 결과 표 원본 (A-1–{A[ORDER[-1]]})'])


def summary():
    s, top = slide('summary', '요약', '한 장 요약',
                   '결론 │ 관측을 늘리고 관측별 점수를 맞춰 결합해, AD2 16-shot에서 SuperAD보다 AUPRO +4.17pp')
    rows = [
        ['문제', '소수 정상 support에서 관측을 어떻게 늘리고 합칠까?', '큰 거리의 원인이 모호하고, 관측을 늘리면 단위·정상 반응이 섞임',
         f'{P["ambig"]}–{P["hard2"]}'],
        ['방법', '무엇을 바꿨나?', '관측별 bank + 정상 q95 단위 정렬 + Shallow 고정 가중(λ = 0.125)', f'{P["overview"]}–{P["control"]}'],
        ['주 결과', 'AD2 few-shot 경쟁력은?', '1–16 shot 3개 지표 모두 1위, 16-shot 3-seed 60.06 vs 55.89', f'{P["t4"]}·{P["t4aux"]}'],
        ['구성 분석', '이득은 어디서 오나?', '다해상도 +4.95 / +2.10pp, 단일 bank 대비 +2.09pp, λ = 1 대비 +5.55pp',
         f'{P["t8"]}–{P["t10ab"]}'],
        ['한계', '아직 확인 안 된 것은?', 'q95의 정규화 생략 대비 이득 없음, linear0125 private·shift·비용 미확보',
         f'{P["t8cd"]}·{P["pending"]}'],
    ]
    table(s, .70, top + .05, [1.4, 3.7, 5.64, 1.2], ['섹션', '질문', '결론', '쪽'], rows, size=13, bottom=5.9)


# ---------------------------------------------------------------- 01 introduction
def bg():
    s, top = slide('bg', S1, '배경: 정상 몇 장으로 결함 위치 찾기',
                   '목적 │ 재학습 없이 정상 1–16장만으로 결함 영역을 짚는 few-shot 이상 국소화',
                   ['동결 백본 재사용: 대규모 사전학습 표현(DINOv2)은 별도 학습 없이 여러 과제로 전이됨',
                    '패치 매칭: query 패치와 가장 가까운 정상 패치의 거리로 결함 위치를 찾음 → 제품별 재학습 불필요',
                    '주 지표 AUPRO@0.05: 오경보율 5% 이하에서 결함 영역을 얼마나 덮는가'],
                   note=f'예시: MVTec AD hazelnut. {MAPNOTE}. {SRC} §1.1 (Oquab 2023; Roth 2022; Damm 2025)')
    figure(s, CH / 'task.png', top + .05)


def ambig():
    s, top = slide('ambig', S1, '문제: 큰 거리가 항상 결함은 아니다',
                   '문제 │ 큰 거리는 결함에서도, support에 없던 정상 변동에서도 생겨 원인을 구분할 수 없다',
                   ['정상 support의 한계: 소수 정상 이미지는 조명 드리프트·자세·난반사·센서 노이즈를 다 담지 못함',
                    '일반 vision 근거: 분포 변화 아래 성능 저하(Hendrycks 2019)와 점수 신뢰도 변화(Ovadia 2019)',
                    '결론: "큰 거리"만으로는 결함인지 미관측 정상 변동인지 확정할 수 없음 (High-Score Ambiguity)'],
                   note=f'그림은 개념도(가상 분포)이며 실측이 아님. {SRC} §1.1')
    figure(s, CH / 'ambiguity_concept.png', top + .05)


def expand():
    s, top = slide('expand', S1, '동기: 같은 이미지를 더 많이 관찰한다',
                   '동기 │ 정상 이미지를 늘릴 수 없으면, 같은 이미지를 여러 깊이·해상도에서 관찰해 단서를 늘린다',
                   ['표현 계층: Deep(L11·23)은 안정적 의미 기준선, Shallow(L5)는 미세 결함에 민감하나 정상 외관에도 반응',
                    '입력 해상도: 448 px 토큰 1개 = 짧은 변 3.1% → 1120 px = 1.25%, 미세 결함의 공간 표본 밀도 증가',
                    '근거: 다계층(DPT, Raghu 2021)·다중 척도(FPN, S², FeatUp)가 dense prediction을 개선'],
                   note=f'아래: hazelnut query의 해상도별 ScopeAD 맵. {MAPNOTE}. {SRC} §1.2, §2.2')
    figure(s, CH / 'resolution_row.png', top + .1)


def hard1():
    s, top = slide('hard1', S1, '결합의 과제 ①: bank를 하나로 합치면',
                   '과제 ① │ 계층 특징을 한 bank로 합치면 관측별 정상 변동이 한 거리 안에 섞인다',
                   ['기존 패치 매칭: 여러 계층 특징을 이어 붙여 bank 1개, 점수 1개 (PatchCore 등)',
                    '근거: 같은 저장 용량에서 단일 bank는 AUPRO@0.05 57.48%, 관측별 bank는 59.57% (표 8)',
                    '대가: bank를 나누면 bank마다 정상 거리 분포가 달라 → 결합 전에 점수 단위를 맞춰야 함'],
                   note=f'개념도. 수치는 AD2 test_public 16-shot seed 0. {SRC} §1.2 "관측 결합의 어려움", 표 8')
    figure(s, CH / 'bank_schematic.png', top + .2)


def hard2():
    s, top = slide('hard2', S1, '결합의 과제 ②: Shallow를 그대로 더하면',
                   '과제 ② │ 단위를 맞춰도 Shallow 점수에는 정상 질감 반응이 섞여, 그대로 더하면 성능이 떨어진다',
                   ['그림: λ = 1이면 정상 walnut 전체가 밝아짐(최댓값 1.22 → 2.52), λ = 0.125는 Deep만과 비슷',
                    '근거: q95 정렬 뒤 Shallow 전체 가산(λ = 1) 54.02% vs 고정 가중(λ = 0.125) 59.57% (표 8b)',
                    '의미: 척도 보정만으로는 Shallow 기여량 제어를 대신할 수 없음'],
                   note=f'AD2 walnuts 정상·결함 query, 하늘색 = GT(표시용). {MAPNOTE}. 수치 근거: 표 8b (AD2 16-shot seed 0). {SRC} §1.2')
    figure(s, CH / 'shallow_leak.png', top + .05)


def split():
    s, top = slide('split', S1, '해법: 단위와 비중을 따로 정한다',
                   '해법 │ 결합을 "단위 맞추기(Scale Calibration)"와 "비중 정하기(Contribution Control)"로 나눈다',
                   ['Scale Calibration: 관측별 bank의 정상 LSO 거리 q95를 1로 — 레이블 없이 정상 이미지만으로 정한 단위',
                    'Contribution Control: Deep 평균 기준선에 Shallow를 λ = 0.125로 가산 (AD2 2-shot에서 선택 후 동결)',
                    '근거: q95 정렬은 MAD 정렬보다 AUPRO@0.05 +0.76 ~ +1.69pp (모든 shot, 표 8d)'],
                   note=f'점수 정규화·이질적 점수 결합의 동기: Kim 2021(CAM), Wang 2022(ViM); 레이블 기반 calibration(Guo 2017)과 구분. {SRC} §1.2')
    figure(s, CH / 'formula.png', top + .25, bottom=min(6.6, top + 3.0))


def roles():
    s, top = slide('roles', S1, '세 요소의 역할 분담',
                   '구조 │ 관측 확장과 두 결합 규칙이 서로 다른 단계를 맡아, 한 요소가 다른 요소를 대신하지 않는다',
                   note=f'AD2 test_public 16-shot, seed 0, Full = 59.57%. 단일 seed 개발 비교이며 개별 차이를 합산 기여로 해석하지 않음. '
                        f'{SRC} §1.2 "관측과 결합의 역할 분담", 표 8')
    rows = [
        ['다해상도 관측', '관측', '같은 이미지를 448–1120 px 4개 해상도로 관찰', '448 단일 해상도: −7.56pp'],
        ['관측별 bank', '결합', '(층, 해상도, bag)마다 bank와 점수 맵을 따로 유지', '계층 concat 단일 bank: −2.09pp'],
        ['Scale Calibration', '결합', '관측별 정상 q95로 점수 단위 정렬', 'MAD 보정: −1.05pp'],
        ['Contribution Control', '결합', 'Shallow 비중을 λ = 0.125로 고정', '3계층 동일 가중 평균: −1.36pp'],
    ]
    table(s, .70, top + .15, [2.6, 1.2, 4.6, 3.54], ['요소', '단계', '하는 일', '이것만 바꾸면 (AUPRO@0.05)'], rows,
          size=13, bottom=5.0)


def contrib():
    s, top = slide('contrib', S1, '주요 기여',
                   '기여 │ 관측 확장, 관측별 bank의 보정·제어 결합, 그리고 두 이득을 통제 실험으로 분리 검증',
                   note=f'모두 개발에 사용한 모집단(AD2 public, MVTec AD, VisA) 결과. "SOTA"는 Outline §1.3 표현. {SRC} §1.3')
    rows = [
        ['① Observation\nExpansion', '재학습 없이 해상도(토큰 3.1% → 1.25%)와 계층(Shallow / Deep)\n두 관측 축을 확장',
         '최선 단일 해상도 대비\n+4.95 / +2.10pp (2 / 16-shot)', f'{P["t10ab"]}·{P["t10"]}'],
        ['② Calibrated &\nControlled Fusion', '관측별 bank + q95 척도 보정 + Shallow 고정 가중 가산,\n학습 없는 결정론적 파이프라인',
         '단일 bank 대비 +2.09pp,\nShallow 전체 가산 대비 +5.55pp', f'{P["t8"]}·{P["t8b"]}'],
        ['③ Disentanglement\n& SOTA', '관측 확장(구성) 이득과 결합 이득을 통제 실험으로 분리,\nAD2 few-shot 비교', 'SuperAD 대비 3-seed\nAUPRO@0.05 +4.17pp (16-shot)',
         str(P['t4'])],
    ]
    table(s, .70, top + .15, [2.5, 5.2, 3.24, 1.0], ['기여', '내용', '근거', '쪽'], rows, size=13, bottom=5.6)


def related():
    s, top = slide('related', S1, '관련 연구와 차별점',
                   '위치 │ 같은 최근접 패치 매칭 계열이지만, 관측별 bank와 정상 기준 단위 정렬로 결합을 다르게 한다',
                   note=f'{SRC} §2.1–2.3, 부록 A.4 (서지 검증 2026-10-08).')
    rows = [
        ['Few-shot 패치 매칭', 'PatchCore, AnomalyDINO,\nSuperAD', '계층 특징을 이어 붙인 bank,\nDINOv2 최근접 거리', '관측별 bank 분리 +\n결합 전 정상 기준 단위 정렬'],
        ['다중 척도·다계층 표현', 'FPN, S², FeatUp,\nDPT, Raghu 2021', '다중 척도·계층 표현으로\ndense prediction 개선', '학습 없이 해상도별 정상 거리 맵 평균\n(S²는 타일 재조립·채널 연결)'],
        ['점수 정규화·결합', 'Kim 2021 (CAM),\nViM 2022', '점수 맵 정규화, 이질적 점수의\n척도를 맞춘 뒤 결합', '정상 support LSO 거리의 q95로\n단위 결정 (레이블 불필요)'],
        ['분포 변화·저-FPR 평가', 'ImageNet-C, Guo 2017,\nOvadia 2019, MVTec AD 2', 'shift 아래 성능·신뢰도 평가,\n조명 변화 산업 벤치마크', '저-FPR 곡선 지표와\n고정 임계값 결과를 구분해 평가'],
    ]
    table(s, .70, top + .1, [2.6, 2.8, 3.1, 3.44], ['연구 흐름', '대표 연구', '한 일', 'ScopeAD와의 차이'], rows,
          size=12, bottom=6.4)


# ---------------------------------------------------------------- 02 method
def overview():
    s, top = slide('overview', S2, 'ScopeAD 전체 흐름',
                   '구조 │ 정상 support로 48개 bank와 단위 c_m을 한 번 구하고(A), query마다 정렬 → 가산 → 평균',
                   notes='A 피팅: 층 3 × 해상도 4 × bag 4 = 48개 bank, 정상 LSO 거리의 q95로 c_m (1회). '
                         'B–D 추론: z = d / c → S = z_d + 0.125·z5 → bag·4개 해상도 평균 → 이미지 점수 = 상위 1% 픽셀 평균.',
                   note='그림: PaperWorks/figures/draw_scarf_overview.py를 linear0125·ScopeAD 표기로 고친 발표용 버전(CQ 미표시). '
                        f'히트맵은 캐시된 walnuts 계층 맵 재합성. 원고 그림 1은 재작도 필요(Outline A.3). {SRC} §3.1')
    figure(s, CH / 'overview_scopead.png', top + .02)


def bank():
    s, top = slide('bank', S2, '관측 구성: 4개 해상도 × 4개 bag',
                   '설계 │ 해상도마다 토큰 밀도가 다르고, 이미지당 256토큰을 4번 뽑아 bag 평균으로 표본 우연성을 줄인다',
                   ['Bag: 이미지당 256토큰 비복원 추출 × 시드 4개, 세 층은 같은 위치 토큰 공유 → 16-shot bank = 4,096토큰',
                    '효과: bag 간 픽셀 표준편차 평균 0.038, 4-bag 평균으로 희소 표본화 차이를 완화 (아래 그림 M1)'],
                   note=f'그림 M1: hazelnut 1120 px, 이미지당 256토큰, 오른쪽 = 픽셀별 bag 간 표준편차. {SRC} §3.2')
    rows = [['448', '32', '약 3.1%', '가장 성긴 관측, 거시 형태'], ['672', '48', '약 2.1%', '중간 수준의 구조 정보'],
            ['896', '64', '약 1.6%', '국소 결함 표본화 밀도 강화'], ['1120', '80', '약 1.25%', '가장 촘촘한 관측, 미세 결함']]
    table(s, .70, top + .05, [2.2, 2.6, 3.0, 4.14], ['해상도 (px)', '짧은 변 토큰 수', '토큰 1개가 덮는 길이', '관측 특성'],
          rows, size=12, row_h=.3)
    figure(s, CH / 'fig_rsg4_mech_bagging_nocap.png', top + .05 + .42 + .3 * 4 + .12)


def scale():
    s, top = slide('scale', S2, 'Scale Calibration: 관측 간 단위 정렬',
                   '방법 │ 각 bank의 정상 LSO 거리 q95(c_m)로 나눠 z = 1을 관측 공통의 기준점으로 맞춘다',
                   ['같은 거리, 다른 뜻: d = 0.178은 Layer 5 기준(0.123)은 넘고 Layer 23 기준(0.232)은 못 넘음 (오른쪽)',
                    'LSO: support 이미지 하나를 빼고 나머지와 매칭한 정상 거리(16-shot 4,096개)로 기준을 정함',
                    '범위: z = 1은 결합 전 관측 간 비중을 맞추는 기준이며, 최종 판정 임계값이 아님'],
                   note=f'c_m = Q0.95(D_m^LSO), z_m = d_m / c_m. 그림 M2: hazelnut 1120 px, bag 평균, d*는 예시값. {SRC} §3.3')
    figure(s, CH / 'fig_rsg4_mech_scale_nocap.png', top + .02)


def control():
    s, top = slide('control', S2, 'Contribution Control: Shallow 고정 가중 가산',
                   '방법 │ S = z_d + 0.125·z5: Deep 평균을 기준선으로 두고, Shallow는 작은 고정 가중치로만 더한다',
                   ['λ = 0.125: AD2 2-shot seed 0 AUPRO로 {0, 0.125, 0.25, 0.5, 1, 2} 중 선택해 동결 (post-hoc)',
                    '예시: 전체 가산은 정상 질감 면을 1.90까지 올리지만, ScopeAD는 0.806에 머묾',
                    '게이트·임계값·학습 가중치 없음, 계층마다 bank와 점수 맵을 유지한 채 점수 단위에서 결합'],
                   note=f'z_d = (z11 + z23)/2. 막대 값은 설명용 예시(실측 아님), 전체 가산 = z_d + z5. 다른 데이터셋에서 λ 재조정 없음. {SRC} §3.4')
    figure(s, CH / 'fusion_example.png', top + .05)


def maps():
    s, top = slide('maps', S2, '실제 맵에서 본 Shallow 가산',
                   '확인 │ 0.125·z5의 가산분은 작아 Deep 맵을 크게 바꾸지 않고, λ = 1처럼 정상 껍질을 밝히지 않는다',
                   ['Shallow z5 전체: 균열 테두리와 함께 정상 껍질 윤곽에도 반응',
                    '0.125·z5: 같은 색 범위에서 거의 어둡게 보일 만큼 작은 가산분',
                    'λ = 1 비교: 정상 껍질까지 밝아짐 → 과제 ②의 정상 반응 누적과 같은 현상'],
                   note=f'MVTec AD hazelnut, 1120 px, bag 평균, 색 범위 [0, 2], 하늘색 = GT(표시용). {MAPNOTE}. '
                        f'원고의 그림 M3는 linear0125로 재작도 예정(Outline §3.4). {SRC} §3.4')
    figure(s, CH / 'fusion_maps.png', top + .15)


def spec():
    s, top = slide('spec', S2, '맵 통합과 고정 사양',
                   '사양 │ bag 평균 → 4개 해상도 보간 평균으로 최종 맵을 만들고, 모든 값은 사전에 고정한다',
                   note=f'{SRC} §3.5–3.6, 표 8 (CQ99는 표 8 실행 조건).')
    rows = [
        ['백본', 'DINOv2 ViT-L/14 (동결, 약 304M, 패치 14)', ''],
        ['층 / 해상도 / Bag', 'L5 (Shallow), L11·23 (Deep) / 448·672·896·1120 / 4 × 256토큰', '모든 query에 고정'],
        ['척도 분위수 q', '0.95 (정상 LSO 거리)', 'Scale Calibration'],
        ['Shallow 가중치 λ', '0.125 (Layer 5에만)', 'AD2 2-shot seed 0로 선택, 이후 재조정 없음'],
        ['맵 통합', '해상도별 맵을 원본 크기로 bilinear 보간 → 동일 가중 평균', ''],
        ['이미지 점수', '최종 맵 상위 1% 픽셀 평균', 'SuperAD·PatchCore I-AUROC 집계와 동일'],
        ['정밀도 · 용량', 'FP16 autocast, bank 256토큰 / source', ''],
        ['재현성', 'seed 0·1·2, 고정 support manifest', '역전파·파라미터 갱신·테스트 라벨 없음'],
    ]
    table(s, .70, top + .1, [2.6, 5.6, 3.74], ['항목', '값', '비고'], rows, size=12, bottom=6.5)


# ---------------------------------------------------------------- 03 results
def setup():
    s, top = slide('setup', S3, '평가 설정 (표 1·2)',
                   '설정 │ 주 비교는 조명 변화가 섞인 MVTec AD 2 test_public 8클래스, 정상 1–16장',
                   ['Support: regular 조명 train/good에서 선택 → Shift 1–3 조명은 support에 없던 변화',
                    'Seed: 비교군은 seed 0–2 평균 ± SD, ScopeAD는 2·16-shot 3-seed, 1·4·8-shot seed 0 단일',
                    '지표: AUPRO@0.05(주), pAUROC@0.05, I-AUROC(맵 상위 1% 평균) — 모두 임계값 없이 계산'],
                   note=f'{SRC} §4.1 표 1–2. 공식 private F1 임계값: support 제외 train/good 8번째 간격 샘플의 pixel p95 × 1.421 (private GT 미사용).')
    rows = [[pretty(c) for c in r] for r in otable('**표 1.')[1:]]
    table(s, .70, top + .1, [2.6, 3.6, 5.74], ['Split', '이미지 수', '조명 · 레이블 범위'], rows, size=12, bottom=6.0)


def baselines():
    s, top = slide('baselines', S3, '비교군 설정 (표 3)',
                   '설정 │ 비교군은 유효 입력을 1120으로 바꿔 다시 실행하고, 백본·계층·후처리는 각 방법 설정을 유지한다',
                   note='같은 support 경로·query·evaluator 사용. RegAD는 공식 STN이 224 고정. FastRecon·RAD는 논문 식의 독립 구현. '
                        f'SuperAD는 공통 support 사용(원 논문은 전체 train coreset). 외부 비교 = 전체 파이프라인 비교. {SRC} §4.1 표 3')
    rows = [[pretty(c) for c in r] for r in otable('**표 3.')[1:]]
    rows = [[r[0].replace(' (paper configuration)', ' (paper)'), r[1], r[2], r[3]] for r in rows]
    table(s, .70, top + .05, [2.4, 3.2, 2.9, 3.44], ['방법', '학습 · support fit', 'Backbone / layers', '입력 · 해상도'],
          rows, size=10, bottom=6.72)


def t4():
    s, top = slide('t4', S3, '표 4 · Few-shot 주 비교: AUPRO@0.05',
                   '결과 │ 1–16 shot 모두 1위, 16-shot 3-seed 60.06 ± 0.88 vs SuperAD 55.89 ± 1.07 (+4.17pp)',
                   ['격차: 비교군 최고값보다 +3.24 ~ +8.50pp, 정상 이미지가 가장 적은 1-shot에서 가장 큼',
                    '통계 주의: ScopeAD 2·16-shot은 3-seed, 1·4·8-shot은 seed 0 단일값 (비교군은 모두 3-seed)'],
                   note='AD2 test_public 8-class macro. 비교군 1120 입력 FP32, ScopeAD FP16·4 해상도. RegAD는 224(공식 모델). '
                        f'표 원본 {A["a_t4"]}. {SRC} §4.2 표 4')
    figure(s, CH / 't4_aupro.png', top + .05)


def t4aux():
    s, top = slide('t4aux', S3, '표 4 · 보조 지표: pAUROC@0.05 · I-AUROC',
                   '결과 │ pAUROC@0.05와 I-AUROC도 모든 shot에서 1위 (16-shot 77.61 / 75.92)',
                   ['pAUROC@0.05: 16-shot SuperAD 74.60 대비 +3.01pp, 1-shot부터 비교군 최고값을 넘음',
                    'I-AUROC: 16-shot 비교군 최고(AnomalyDINO 72.07) 대비 +3.85pp'],
                   note=f'굵은 선 = ScopeAD, 진한 회색 = SuperAD, 연한 회색 = 나머지 비교군. 표 원본 {A["a_t4aux"]}. {SRC} §4.2 표 4')
    figure(s, CH / 't4_aux.png', top + .3)


def t5():
    s, top = slide('t5', S3, '표 5 · MVTec AD · VisA few-shot',
                   '결과 │ I-AUROC는 10개 조건 중 7개 1위, P-AUROC는 최고값과 0.04–1.36pp 이내',
                   ['I-AUROC 2위인 조건: MVTec AD 16-shot, VisA 1·2-shot (모두 SuperAD 다음)',
                    'P-AUROC: MVTec AD 2-shot·VisA 16-shot 1위, 최대 차이는 VisA 1-shot SubspaceAD 98.09'],
                   note='seed 0, MVTec AD 15·VisA 12 클래스 macro, support hash 일치(E9 매니페스트). AnomalyDINO·RAD·SubspaceAD는 1120 입력. '
                        f'FastRecon = DINOv2 기반 프로젝트 변형. linear0125 hash 재감사는 수집 단계 과제. 표 원본 {A["a_t5"]}. {SRC} §4.2 표 5')
    figure(s, CH / 't5.png', top + .3)


def t5b():
    s, top = slide('t5b', S3, '표 5b · MVTec AD · VisA 3-seed',
                   '확인 │ 3-seed 평균도 seed 0와 비슷한 수준을 유지한다 (16-shot I-AUROC 98.52 / 95.78)',
                   ['seed 0 대비: MVTec AD 16-shot I-AUROC 98.74 → 3-seed 98.52, VisA 95.99 → 95.78',
                    '주의: seed 1·2의 support는 비교군과 짝지어지지 않아 표 5의 비교군과 직접 순위를 매기지 않음'],
                   note=f'mean ± sample SD, 각 셀 I-AUROC / P-AUROC (%). {SRC} §4.2 표 5b')
    rows = [[pretty(c) for c in r] for r in otable('**표 5b.')[1:]]
    table(s, 1.6, top + .3, [2.4, 3.9, 3.9], ['Dataset', '2-shot', '16-shot'], rows, size=14, bottom=top + 1.9)


def t6():
    s, top = slide('t6', S3, '표 6 · MVTec AD 2 공식 private',
                   '참고 │ 공식 private은 아직 이전 게이트 설정 결과이며, linear0125로 재제출해 교체해야 한다',
                   ['게이트 설정 결과: AU-PRO 64.80 / 62.49 — full-data 방법과 같은 범위 (데이터 예산이 달라 순위화 안 함)',
                    'SegF1 34.84 / 30.68: 고정 임계값(train/good p95 × 1.421) 성능은 낮음 → 임계값 설계가 과제'],
                   note='● TESTpriv, ○ TESTpriv,mix. † Layer 5 β = 0.99 초과 잔차를 더한 이전 설정의 제출 결과(Outline 표 6 캡션). '
                        f'Full-data 행은 원 논문·리더보드 수치. 전체 표 {A["a_t6"]}. {SRC} §4.3.1 표 6')
    figure(s, CH / 't6_private.png', top + .3)


def shift():
    s, top = slide('shift', S3, '그림 2 · 표 7 · 분포 변화 평가',
                   '현황 │ 조명 변화(그림 2)와 합성 손상(표 7)은 linear0125로 아직 평가하지 않았다',
                   ['그림 2: 기존 그림의 SCARF 선은 게이트 결과라 사용하지 않음 → strata 집계를 켜고 재평가',
                    '표 7: 비교군 결과만 확보 — Gaussian noise에서는 PatchCore 74.03이 가장 높음'],
                   note='아래: 표 7 비교군 AUPRO@0.05. MVTec AD 15클래스, 16-shot, seed 0, severity 3, support는 깨끗한 이미지. '
                        f'† DINOv2 기반 프로젝트 변형. 표 원본 {A["a_t7"]}. {SRC} §4.3.2–4.3.3')
    figure(s, CH / 't7_corruption.png', top + .35)


# ---------------------------------------------------------------- 04 ablation
def t8():
    s, top = slide('t8', S4, '표 8 · 구성 요소 소거 (AD2 16-shot)',
                   '결과 │ 어느 요소를 하나만 바꿔도 Full(59.57)보다 낮고, 해상도와 bank 구조의 영향이 가장 크다',
                   ['관측: 448 단일 해상도 −7.56pp로 다해상도 관측의 영향이 가장 큼',
                    '결합: 단일 bank −2.09pp, 3계층 동일 가중 −1.36pp, MAD 보정 −1.05pp',
                    '정규화 생략은 −0.40pp지만 pAUROC@0.05는 +0.19pp → q95의 정규화 생략 대비 이득은 주장하지 않음'],
                   note='동결 DINOv2 ViT-L/14, FP16, 4 해상도, cap256, 4 bags, CQ99, 동일 support. 각 행은 Full에서 한 구성만 바꿈. '
                        f'단일 seed 개발 비교. 표 원본 {A["a_t8"]}. {SRC} §4.4.1 표 8')
    figure(s, CH / 't8_ablation.png', top + .05)


def t8a():
    s, top = slide('t8a', S4, '표 8a · 데이터셋·shot별 구성 소거',
                   '결과 │ 세 데이터셋 모든 shot에서 Full이 3계층 동일 가중 평균과 계층 concat보다 높다',
                   ['계층 concat 대비 (16-shot): AD2 AUPRO +2.09, MVTec AD I-AUROC +0.55, VisA I-AUROC +1.90pp',
                    '정규화 생략: MVTec AD·VisA I-AUROC는 0.07–1.06pp 높고, AD2 AUPRO는 −0.40 ~ +0.33pp로 엇갈림'],
                   note='seed 0 단일값. AD2 1/4/8-shot은 로컬 4090, 2/16-shot Full은 fusion sensitivity 실행의 동일 설정. '
                        f'P-AUROC 포함 표 원본 {A["a_t8a"]}. {SRC} §4.4.1 표 8a')
    figure(s, CH / 't8a.png', top + .3)


def t8b():
    s, top = slide('t8b', S4, '표 8b · Shallow 가중치 λ 민감도',
                   '결과 │ λ ≥ 0.5에서 AUPRO@0.05가 빠르게 낮아지고, λ = 0.125가 2·16-shot 모두 가장 높다',
                   ['16-shot: λ = 0.125 59.57 → λ = 1 54.02 (−5.55pp) → λ = 2 47.33',
                    '선택 절차: 2-shot AUPRO만으로 λ를 고른 뒤 16-shot 해석 전에 동결 (개발 평가 사용, post-hoc)',
                    '해석: 가중치가 얕은 층의 정상 질감 반응이 결합 점수에 섞이는 정도를 정한다는 해석과 일치'],
                   note=f'd = (z11+z23)/2, s = z5. FP16, 4-scale, cap256, q95, CQ99, seed 0. 표에는 양의 가중치 조건만 제시. {SRC} §4.4.2 표 8b')
    figure(s, CH / 't8b_lambda.png', top + .05)


def t8cd():
    s, top = slide('t8cd', S4, '표 8c · 8d · 정규화 방식 비교',
                   '결과 │ q95는 MAD보다 모든 shot에서 높지만, 정규화 생략과는 차이가 작고 방향이 엇갈린다',
                   ['AD2: q95 − MAD AUPRO +0.76 ~ +1.69pp, q95 − none −0.33 ~ +0.40pp (평균 +0.02)',
                    'MVTec AD·VisA: q95가 MAD보다 I-AUROC +0.34 ~ +1.60pp, none보다는 0.07 ~ 1.06pp 낮음',
                    '주장 범위: 단위 정렬은 관측별 bank의 설계 요구, 성능 근거는 "MAD 대비 우위"로만 서술'],
                   note=f'전체 방법(linear0125)에서 normalization만 바꾼 대조, 4 해상도, cap256, 같은 support, seed 0. 표 원본 {A["a_t8cd"]}. '
                        f'{SRC} §4.4.1 표 8c·8d')
    figure(s, CH / 't8cd_norm.png', top + .3)


def t10ab():
    s, top = slide('t10ab', S4, '표 10a · 10b · 단일 vs 다중 해상도',
                   '결과 │ 4개 해상도 결합은 최선의 단일 해상도보다 AUPRO@0.05 +4.95pp(2-shot), +2.10pp(16-shot)',
                   ['단일 해상도는 높을수록 항상 좋지 않음: 2-shot 최선 = 672, 16-shot 최선 = 1120',
                    'MVTec AD·VisA: 1120 단일보다 I-AUROC +0.53 ~ +5.27pp, P-AUROC도 네 조건 모두 상승',
                    '주의: bank당 cap은 같지만 총 용량은 달라 총 용량 효과는 이 표로 분리하지 않음 (→ 표 10)'],
                   note=f'linear0125, q95, FP16, cap256, seed 0. 1120 대비 2-shot 이득은 +5.06pp. 표 원본 {A["a_t10"]}. {SRC} §4.4.3 표 10a·10b')
    figure(s, CH / 't10ab_resolution.png', top + .05)


def t10():
    s, top = slide('t10', S4, '표 10 · 총 bank 용량 통제 비교',
                   '참고 │ 총 토큰 수를 맞춰도 다해상도가 1120 단일보다 +2.52pp — 단, 이전 게이트 설정의 결과',
                   ['용량 통제: 1120 단일에 토큰을 4배(1024) 줘도 56.63%로 오히려 −0.33pp',
                    '주의: 이전 비대칭 게이트 설정(E12) 결과 → linear0125로 같은 비교를 다시 실행해 교체 예정'],
                   note=f'AD2 test_public, 16-shot, seed 0, full8, DINOv2 ViT-L/14. E12 resolution control. 표 원본 {A["a_t10"]}. {SRC} §4.4.3 표 10')
    figure(s, CH / 't10_gate.png', top + .35)


def t11():
    s, top = slide('t11', S4, '표 11 · 척도 분위수 q 민감도',
                   '결과 │ q = 0.80 ~ 0.999에서 AUPRO@0.05 변화는 −0.22 ~ +0.04pp로 평탄하다',
                   ['중앙값 기준(q = 0.5)만 낮음 (−1.11 / −0.55pp), 꼬리확률 변환·토큰별 로컬 척도도 개선 없음',
                    '척도 진단: 컴포넌트 간 q95 척도 차이 평균 약 1.8배(deep끼리 약 1.4배) → 정렬이 비중을 완만하게 조정',
                    'LSO vs LOO-token: AD2 AUPRO 차이 0.12pp 이내 → LSO 선택 자체의 우위는 주장하지 않음'],
                   note=f'linear0125, λ = 0.125 고정, seed 0. 변형 행은 표 원본 {A["a_t11"]}. {SRC} §4.4.1 LSO 대조, §4.4.4 표 11·척도 진단')
    figure(s, CH / 't11_q.png', top + .05)


# ---------------------------------------------------------------- 05 discussion
def discussion():
    s, top = slide('discussion', S5, '논의: 관측 확장과 관측별 결합의 역할',
                   '해석 │ 다해상도와 관측별 bank는 다른 단계에서 기여하고, Shallow 가산은 가중치에 민감하다',
                   [f'관측 확장: 최선 단일 해상도 대비 +4.95 / +2.10pp, 용량 통제(게이트 설정)에서도 +2.52pp ({P["t10ab"]}·{P["t10"]}쪽)',
                    f'관측별 결합: 세 데이터셋 모든 shot에서 단일 bank보다 높음 ({P["t8a"]}쪽)',
                    f'q95 정렬: MAD보다 높지만 정규화 생략과 차이 작음 → 척도 차이가 큰 관측·조건 변화 query에서 확인 ({P["t8cd"]}쪽)',
                    '배포: 단위 정렬이 query FPR을 정하지 않음 → 조명 조건별 임계값 보정과 검증이 필요'],
                   note=f'{SRC} §5.1–5.2, 부록 A.1.')
    boxed(s, .70, top + .3, 11.94, 1.2, '주장 범위',
          '지지(개발 모집단): 관측별 bank > 단일 bank, q95 > MAD, Shallow 기여량 제어 필요, 다해상도 이득, AD2 2·16-shot SuperAD 대비 우위. '
          '미검증: q95의 정규화 생략 대비 이득, linear0125의 private·shift·corruption·계산 비용.')


def pending():
    s, top = slide('pending', S5, '남은 검증',
                   '계획 │ 게이트 설정으로 남은 표 6·10을 교체하고, 분포 변화·비용·다중 seed를 채운다',
                   note=f'{SRC} 부록 A.1–A.3, §4.4.4 후속 실험 상태, §4.5.')
    rows = [
        ['공식 private (표 6)', '게이트 설정 결과 64.80 / 62.49', 'linear0125, 같은 support·임계값 절차로 재제출'],
        ['용량 통제 해상도 비교 (표 10)', '게이트 설정 결과', 'multi-scale vs 1120 (cap256 / 1024) 재실행'],
        ['조명 strata (그림 2)', 'ScopeAD 선 미사용', 'condition strata 집계를 켜고 재평가'],
        ['합성 손상 (표 7)', 'ScopeAD 미실행', 'MVTec AD 15클래스 × 5조건 실행'],
        ['계산 비용 (표 12 · 12a)', 'ScopeAD 미계측', 'linear0125 파이프라인 재계측 (구성 요소별 시간 포함)'],
        ['다중 seed', '2·16-shot 3-seed 완료', '1·4·8-shot seed 1·2'],
        ['local-center 보정', '진행 중 (AD2 2·16-shot)', 'q95·none을 모두 넘는지 확인 후 판단'],
        ['원고 그림 1 · M3 · 3', '게이트 표기 포함 / 사양만 확정', 'linear0125 기준 재작도 (발표용은 캐시 재합성으로 대체)'],
    ]
    table(s, .70, top + .1, [3.4, 3.6, 4.94], ['항목', '현재 상태', '다음 단계'], rows, size=12, bottom=6.5)


def conclusion():
    s, top = slide('conclusion', S5, '결론 및 향후 진행 방향',
                   '결론 │ 관측을 늘리고 관측별 점수를 맞춰 결합하는 ScopeAD로 AD2 few-shot 국소화 1위를 얻었다',
                   note=f'{SRC} §6, 부록 A.1. 근거 쪽: 관측 확장 {P["t10ab"]}, 결합 {P["t8"]}·{P["t8cd"]}, 성능 {P["t4"]}, 남은 항목 {P["pending"]}.')
    rows = [
        ['관측 확장', '4개 해상도 + Shallow / Deep 관측,\n최선 단일 대비 +4.95 / +2.10pp', '용량 통제 비교는 게이트 결과', '표 10 linear0125 재실행'],
        ['관측별 결합', '관측별 bank + q95 정렬 + λ = 0.125,\n단일 bank 대비 +2.09pp', 'q95의 정규화 생략 대비 이득 없음', 'local-center 등 정렬 방식 검증'],
        ['성능', 'AD2 16-shot 3-seed 60.06 ± 0.88,\nSuperAD 대비 +4.17pp', '1·4·8-shot은 seed 0 단일값', '다중 seed 완성'],
        ['강건성 · 배포', '— (private은 게이트 설정 결과)', 'linear0125 private·shift·corruption·비용 미확보', '재제출 · 재평가 · 재계측'],
    ]
    table(s, .70, top + .1, [1.8, 4.0, 3.2, 2.94], ['구분', 'Summary', '현재 문제', '향후 진행 방향'], rows, size=13,
          bottom=6.2)


# ---------------------------------------------------------------- appendix (parsed tables)
def _t4_rows(metric):
    rows = otable('**표 4.')
    out, shot = [], None
    for r in rows[1:]:
        shot = r[0] or shot
        if r[1] == metric:
            out.append([shot] + [pretty(c) for c in r[2:]])
    return out


H4 = ['Shot', 'PatchCore', 'AnomalyDINO', 'FastRecon', 'RAD', 'SubspaceAD', 'SuperAD', 'RegAD (224)', 'ScopeAD']
W4 = [.74] + [1.4] * 8


def a_t4():
    s, top = slide('a_t4', SA, '표 4 원본: AUPRO@0.05',
                   f'원본 │ {P["t4"]}쪽 차트의 수치 (AD2 test_public, 8-class macro, %)',
                   note=f'비교군 seed 0–2 mean ± sample SD. ScopeAD 2·16-shot 3-seed, 1·4·8-shot seed 0 단일값(± 없음). {SRC} 표 4')
    table(s, .70, top + .1, W4, H4, _t4_rows('AUPRO@0.05'), size=11, bottom=5.2, strong_cols=(0, 8))


def a_t4aux():
    s, top = slide('a_t4aux', SA, '표 4 원본: pAUROC@0.05 · I-AUROC',
                   f'원본 │ {P["t4aux"]}쪽 차트의 수치 (%)', note=f'통계량 표기는 A-1과 같음. {SRC} 표 4')
    text_box(s, .70, top - .05, 6, .3, 'pAUROC@0.05', 13, NAVY, BOLD, True, name='TableLabel1')
    table(s, .70, top + .27, W4, H4, _t4_rows('pAUROC@0.05'), size=10, row_h=.3, strong_cols=(0, 8))
    y2 = top + .27 + .42 + .3 * 5 + .15
    text_box(s, .70, y2, 6, .3, 'I-AUROC', 13, NAVY, BOLD, True, name='TableLabel2')
    table(s, .70, y2 + .32, W4, H4, _t4_rows('I-AUROC'), size=10, row_h=.3, strong_cols=(0, 8), name='SummaryTable2')


def a_t5():
    s, top = slide('a_t5', SA, '표 5 원본: MVTec AD · VisA',
                   f'원본 │ {P["t5"]}쪽 차트의 수치 (각 셀 I-AUROC / P-AUROC, 전체 FPR, %)',
                   note=f'seed 0, support hash 일치 클래스 macro. — = 검증된 결과 없음. {SRC} 표 5')
    rows = [[f'{r[0]} · {r[1]}'] + [pretty(c) for c in r[2:]] for r in otable('**표 5.')[1:]]
    head = ['Dataset · Shot', 'PatchCore', 'AnomalyDINO', 'FastRecon', 'RAD', 'SubspaceAD', 'SuperAD', 'RegAD', 'ScopeAD']
    table(s, .70, top + .05, [1.58] + [1.295] * 8, head, rows, size=10, bottom=6.72, strong_cols=(0, 8))


def a_t6():
    s, top = slide('a_t6', SA, '표 6 원본: 공식 private',
                   '참고 │ Ours 행은 이전 게이트 설정 결과, few-shot 비교군의 private 결과는 아직 없음',
                   note=f'각 셀 TESTpriv / TESTpriv,mix (%). Full-data 행은 원 논문·리더보드 수치. {SRC} 표 6')
    rows = [[pretty(c) for c in r] for r in otable('**표 6.')[1:] if pretty(r[1])]
    blank = [r[0] for r in rows if r[1] == '—']
    rows = [[f'Few-shot 비교군 ({", ".join(blank)})', '—', '—', '—']] + [r for r in rows if r[1] != '—']
    table(s, .70, top + .05, [4.44, 2.5, 2.5, 2.5], ['방법', 'AU-PRO@0.05', 'ClassF1', 'SegF1'], rows, size=10,
          row_h=.3)


def a_t7():
    s, top = slide('a_t7', SA, '표 7 원본: 합성 손상',
                   f'원본 │ {P["shift"]}쪽 차트의 수치 (각 셀 AUPRO@0.05 / pAUROC@0.05, %)',
                   note=f'MVTec AD 15클래스, 16-shot, seed 0, severity 3. — = 미실행. † DINOv2 기반 프로젝트 변형. {SRC} 표 7')
    rows = [[pretty(c) for c in r] for r in otable('**표 7.')[1:]]
    rows = [[r[0].replace(' (project adaptation)', ''), *r[1:]] for r in rows]
    table(s, .70, top + .1, [2.44] + [1.9] * 5, ['방법', 'Clean', 'Brightness', 'Contrast', 'Defocus blur', 'Gaussian noise'],
          rows, size=11, bottom=6.2)


def a_t8():
    s, top = slide('a_t8', SA, '표 8 · 8b 원본',
                   f'원본 │ {P["t8"]}·{P["t8b"]}쪽 차트의 수치 (AD2 test_public, seed 0, %)',
                   note=f'위: 표 8 (16-shot). 아래: 표 8b (d = deep 평균, s = z5). {SRC} 표 8, 8b')
    rows = [[pretty(c) for c in r] for r in otable('**표 8.')[1:]]
    table(s, .70, top + .02, [2.9, 1.6, 1.5, 1.6, 1.4, 1.4, 1.54],
          ['구성', 'Bank 구조', '정규화', 'Shallow 항', 'AUPRO', 'pAUROC', 'Δ AUPRO'], rows, size=10, row_h=.29)
    rows2 = [[pretty(c) for c in r] for r in otable('**표 8b.')[1:]]
    table(s, .70, top + .02 + .42 + .29 * 6 + .25, [3.94, 2.0, 2.0, 2.0, 2.0],
          ['결합 규칙', '2-shot AUPRO', '2-shot pAUROC', '16-shot AUPRO', '16-shot pAUROC'], rows2, size=10, row_h=.29,
          name='SummaryTable2')


def a_t8a():
    s, top = slide('a_t8a', SA, '표 8a 원본: 데이터셋·shot별 구성 소거',
                   f'원본 │ {P["t8a"]}쪽 차트의 수치 (AD2 AUPRO / pAUROC, MVTec AD·VisA I-AUROC / P-AUROC, %)',
                   note=f'seed 0 단일값. {SRC} 표 8a')
    rows = [[pretty(c) for c in r] for r in otable('**표 8a.')[1:]]
    table(s, .70, top + .02, [1.8, 1.0, 2.3, 2.3, 2.3, 2.24],
          ['Dataset', 'Shot', 'Full (linear0125)', '정규화 생략', '3계층 동일 가중', '계층 concat'], rows, size=10, row_h=.27)


def a_t8cd():
    s, top = slide('a_t8cd', SA, '표 8c · 8d 원본: 정규화 방식',
                   f'원본 │ {P["t8cd"]}쪽 차트의 수치 (seed 0, %)',
                   note=f'위: 표 8c (MVTec AD·VisA, I-AUROC / P-AUROC). 아래: 표 8d (AD2, AUPRO / pAUROC / I-AUROC). {SRC} 표 8c·8d')
    rows = [[pretty(c) for c in r] for r in otable('**표 8c.')[1:]]
    table(s, .70, top + .02, [2.3, 1.2, 2.8, 2.8, 2.84], ['Dataset', 'Shot', 'none', 'MAD', 'q95 (기본)'], rows, size=10,
          row_h=.26)
    rows2 = [[pretty(c) for c in r] for r in otable('**표 8d.')[1:]]
    table(s, .70, top + .02 + .42 + .26 * 10 + .18, [.9, 2.75, 2.75, 2.75, 1.4, 1.39],
          ['Shot', 'none', 'MAD', 'q95 (기본)', 'q95 − none', 'q95 − MAD'], rows2, size=10, row_h=.26,
          name='SummaryTable2')


def a_t10():
    s, top = slide('a_t10', SA, '표 10 · 10a · 10b 원본: 해상도',
                   f'원본 │ {P["t10ab"]}·{P["t10"]}쪽 차트의 수치 (%)',
                   note=f'표 10은 이전 게이트 설정(E12). 표 10a: AD2 AUPRO / pAUROC. 표 10b: I-AUROC / P-AUROC. seed 0. {SRC} 표 10·10a·10b')
    r10 = [[pretty(c) for c in r] for r in otable('**표 10.')[1:]]
    r10 = [[r[0].replace(' (Resolution setting)', ''), r[1], r[2], r[3], r[4], r[5], r[6]] for r in r10]
    table(s, .70, top + .02, [3.04, 1.0, 1.4, 1.2, 1.4, 1.95, 1.95],
          ['표 10 (게이트)', '토큰/bag', 'AUPRO', 'Δ vs 단일', 'pAUROC', 'Regular', 'Shift 1'], r10, size=10, row_h=.29)
    y = top + .02 + .42 + .29 * 3 + .2
    r10a = [[pretty(c) for c in r] for r in otable('**표 10a.')[1:]]
    table(s, .70, y, [2.9, 1.6, 1.6], ['표 10a (AD2)', '2-shot', '16-shot'], r10a, size=10, row_h=.29, name='SummaryTable2')
    r10b = [[pretty(c) for c in r] for r in otable('**표 10b.')[1:]]
    table(s, 7.05, y, [1.2, 1.15, 1.6, 1.6], ['표 10b', '해상도', '2-shot', '16-shot'], r10b, size=10, row_h=.29,
          name='SummaryTable3')


def a_t11():
    s, top = slide('a_t11', SA, '표 11 원본: 척도 분위수 q',
                   f'원본 │ {P["t11"]}쪽 차트의 수치 (AD2 AUPRO / pAUROC, seed 0, %)', note=f'λ = 0.125 고정. {SRC} 표 11')
    rows = [[pretty(c) for c in r] for r in otable('**표 11.')[1:]]
    table(s, .70, top + .05, [4.2, 2.5, 2.5, 2.74], ['척도 분위수 q / 변형', '2-shot', '16-shot', 'Δ AUPRO (2 / 16-shot)'],
          rows, size=11, bottom=6.0)


for key in ORDER:
    globals()[key]()

groups = [('표지 · 요약', ['cover', 'toc', 'summary']),
          ('01 Introduction', ['bg', 'ambig', 'expand', 'hard1', 'hard2', 'split', 'roles', 'contrib', 'related']),
          ('02 Method', ['overview', 'bank', 'scale', 'control', 'maps', 'spec']),
          ('03 실험 결과', ['setup', 'baselines', 't4', 't4aux', 't5', 't5b', 't6', 'shift']),
          ('04 Ablation', ['t8', 't8a', 't8b', 't8cd', 't10ab', 't10', 't11']),
          ('05 논의', ['discussion', 'pending', 'conclusion']),
          ('Appendix A', ORDER[MAIN_LAST:])]
assert [k for _, ks in groups for k in ks] == ORDER
set_sections(prs, [(n, [slides[k] for k in ks]) for n, ks in groups])
prs.save(OUT)
print('saved', OUT, len(prs.slides), 'slides')
