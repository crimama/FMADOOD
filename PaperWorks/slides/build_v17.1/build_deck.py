"""Update the existing DSBA deck from the v17.1 outline; retain the original file.
Run with system python3 (python-pptx). All chart values are embedded and editable.
"""
import hashlib
import json
import re
import shutil
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, '/home/hun/.codex/skills/dsba-slide-deck/scripts')
from deck_style import text_box, table, style_run, NAVY, INK, SUB, FONT, BOLD, XBOLD
from deck_ops import fix_table_ids, set_sections

ROOT = Path(__file__).resolve().parents[3]
BUILD = Path(__file__).resolve().parent
SRC = ROOT / 'PaperWorks/slides/SCARF_Outline_v16.3_발표.pptx'
OUT = ROOT / 'PaperWorks/slides/SCARF_Outline_v17.1_발표.pptx'
MD = ROOT / 'PaperWorks/outlines/Outline_v17.1.md'
md = MD.read_text()
backup = SRC.parent / 'backup' / (SRC.stem + '_before_v17.1_20261012.pptx')
if not backup.exists():
    shutil.copy2(SRC, backup)
prs = Presentation(SRC)
assert len(prs.slides) == 35
colors = [NAVY, 'B45D42', '508C91', '7D729A', '8593A4', 'A69A80', '727A6B', 'B4BBC4']
shots = ['1', '2', '4', '8', '16']
source_footer = '출처: Outline_v17.1.md · 2026-10-12 · linear0125 · 개발 평가'

def clean(s):
    return s.replace('**', '').replace('<u>', '').replace('</u>', '').replace('$', '').strip()

def mtable(n):
    start = md.index('**표 ' + n + '.')
    lines = md[start:].splitlines()
    rows = []
    begun = False
    for line in lines:
        if line.startswith('|'):
            begun = True
            row = [clean(x) for x in line.strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+', x) for x in row):
                rows.append(row)
        elif begun:
            break
    return rows[0], rows[1:]

def num(s):
    return float(re.search(r'[-+]?\d+\.\d+', s).group())

def frame(n, sec, title, lead, detail=(), note='', refs=''):
    s = prs.slides[n-1]
    for sh in list(s.shapes):
        if sh.name != 'Logo':
            el = sh._element
            el.getparent().remove(el)
    text_box(s, .27, .20, 10.8, .33, sec, 14, NAVY, BOLD, True, name='SectionLabel')
    text_box(s, .70, .67, 11.0, .49, title, 28, NAVY, XBOLD, True, name='Title')
    text_box(s, .70, 1.29, 11.94, .40, lead, 16, NAVY, BOLD, True, name='Lead')
    y = 1.83
    for line in detail:
        text_box(s, .80, y, 11.65, .38, '• ' + line, 16, INK, name='Detail')
        y += .43
    text_box(s, .70, 6.83, 11.94, .30, note or source_footer, 9, SUB, name='MetricNote')
    text_box(s, 9.62, 7.10, 3, .30, str(n) if n <= 26 else 'A-' + str(n-26), 10, SUB,
             align=PP_ALIGN.RIGHT, name='Page')
    s.notes_slide.notes_text_frame.text = (
        f'기준 문서: {MD.relative_to(ROOT)}\n슬라이드 {n}: {title}\n'
        f'{lead}\n' + '\n'.join(detail) + '\n' + (note or source_footer) + '\n' + refs +
        '\n모든 결과는 이미 개발에 사용한 모집단의 평가다. 이전 게이트 결과를 현재 방법의 결과로 사용하지 않는다.'
    )
    return s

def rows(s, header, data, widths, y=2.20, size=14, row_h=.48):
    return table(s, .70, y, widths, header, [[str(v) for v in r] for r in data],
                 size=size, row_h=row_h, bottom=6.65)

def statement(s, y, title, body, h=.64):
    text_box(s, .80, y, 2.3, h, title, 19, NAVY, BOLD, True)
    text_box(s, 3.05, y, 9.3, h, body, 18, INK)

def rank_columns(frame, pairs=False):
    """Keep the outline's best/runner-up emphasis after transposing its tables."""
    tbl=frame.table
    for col in range(1,len(tbl.columns)):
        values=[tbl.cell(row,col).text for row in range(1,len(tbl.rows))]
        split=[v.split(' / ') if pairs else [v] for v in values]
        ranks=[]
        for k in range(2 if pairs else 1):
            valid=[num(v[k]) for v in split if re.search(r'\d+\.\d+',v[k])]
            ranks.append(sorted(set(valid),reverse=True)[:2])
        for row,parts in enumerate(split,1):
            tf=tbl.cell(row,col).text_frame
            tf.clear()
            p=tf.paragraphs[0]
            p.alignment=PP_ALIGN.CENTER
            for k,part in enumerate(parts):
                if k:
                    run=p.add_run();run.text=' / ';style_run(run,12)
                value=num(part) if re.search(r'\d+\.\d+',part) else None
                best=value is not None and value==ranks[k][0]
                second=value is not None and len(ranks[k])>1 and value==ranks[k][1]
                run=p.add_run();run.text=part
                style_run(run,12,NAVY if best else INK,BOLD if best else FONT,best)
                run.font.underline=second

def equation(s, formula, y=3.15, size=30):
    text_box(s, 1.0, y, 11.3, .8, formula, size, NAVY, BOLD, True, align=PP_ALIGN.CENTER)

def chart(s, cats, series, y=2.55, h=3.85, low=None, high=None, ylabel='', kind='line', x=.80, w=11.7):
    data = CategoryChartData()
    data.categories = cats
    for name, vals in series:
        data.add_series(name, vals)
    typ = XL_CHART_TYPE.LINE_MARKERS if kind == 'line' else XL_CHART_TYPE.COLUMN_CLUSTERED
    ch = s.shapes.add_chart(typ, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.name = FONT
    ch.legend.font.size = Pt(12)
    ch.chart_style = 13
    for axis in [ch.category_axis, ch.value_axis]:
        axis.tick_labels.font.name = FONT
        axis.tick_labels.font.size = Pt(12)
        axis.tick_labels.font.color.rgb = RGBColor.from_string(SUB)
    if low is not None:
        ch.value_axis.minimum_scale = low
    if high is not None:
        ch.value_axis.maximum_scale = high
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string('E4E8EC')
    ch.value_axis.has_title = bool(ylabel)
    if ylabel:
        ch.value_axis.axis_title.text_frame.text = ylabel
        for p in ch.value_axis.axis_title.text_frame.paragraphs:
            for r in p.runs:
                r.font.name = FONT
                r.font.size = Pt(12)
    for i, sr in enumerate(ch.series):
        color = colors[i % len(colors)]
        sr.format.line.color.rgb = RGBColor.from_string(color)
        sr.format.line.width = Pt(2.5 if i == 0 else 1.75)
        sr.format.fill.solid()
        sr.format.fill.fore_color.rgb = RGBColor.from_string(color)
        if kind == 'line':
            sr.marker.style = XL_MARKER_STYLE.CIRCLE
            sr.marker.size = 6
            sr.marker.format.fill.solid()
            sr.marker.format.fill.fore_color.rgb = RGBColor.from_string(color)
            sr.marker.format.line.color.rgb = RGBColor.from_string(color)
    return ch

S1, S2, S3, S4, S5 = '01. 문제와 설계', '02. 방법', '03. 실험 결과', '04. 소거와 보정 분석', '05. 결론과 다음 검증'
SA = 'Appendix A. 상세 결과와 설정'
h4, t4 = mtable('4')
h5, t5 = mtable('5')
h8a, t8a = mtable('8a')
h8d, t8d = mtable('8d')
h11, t11 = mtable('11')

# 1. Cover uses the original deck's logo and template.
s = prs.slides[0]
for sh in list(s.shapes):
    if sh.name != 'Logo':
        sh._element.getparent().remove(sh._element)
text_box(s, .9, 2.25, 11.5, .8, 'SCARF', 44, NAVY, XBOLD, True, name='Title')
text_box(s, .9, 3.10, 11.5, .55, '관측별 Memory Bank와 Support 기반 점수 정렬', 25, NAVY, BOLD, True)
text_box(s, .9, 3.85, 11.5, .55, 'Few-shot 산업 이상 국소화 · linear0125 설정', 20, INK)
line = s.shapes.add_connector(1, Inches(.9), Inches(4.55), Inches(12.4), Inches(4.55))
line.line.color.rgb = RGBColor.from_string(NAVY)
line.line.width = Pt(1)
text_box(s, .9, 4.8, 11.5, .4, '연구 아웃라인 v17.1 발표  ·  2026-10-12', 16, SUB)
text_box(s, .9, 5.45, 11.5, .4, 'AD2 중심 성능 검증과 score calibration 후속 분석', 17, INK)
s.notes_slide.notes_text_frame.text = 'Outline_v17.1.md 기준. 기존 SCARF 약어의 Asymmetric 표현은 방법 변경 후 재검토 대상이므로 영문 확장명을 표지에 사용하지 않았다.'

s = frame(2, '목차', '관측 확장에서 성능 검증까지', '흐름 │ 설계 → 전체 방법 → 주 비교 → 보정의 한계 → 다음 검증')
for y, a, b in [(2.2,'01 문제와 설계','같은 이미지에서 공간·계층 관측을 늘리는 이유'),
               (3.0,'02 방법','별도 bank → support q95 → 고정 선형 결합'),
               (3.8,'03 실험 결과','AD2 주 비교 · MVTec AD/VisA · 3-seed'),
               (4.6,'04 소거와 분석','Bank 구조 · 해상도 · calibration'),
               (5.4,'05 다음 검증','보정 기여 확보와 아직 남은 평가')]:
    statement(s,y,a,b)

s = frame(3, '요약', '전체 방법은 강하고, 보정 기여는 남아 있다', '요약 │ AD2 성능과 관측별 bank의 근거를 중심으로 정리한다')
rows(s, ['질문','현재 근거','해석'], [
 ['전체 성능','AD2 16-shot: 60.06 ± 0.88','3-seed AUPRO@0.05'],
 ['Bank 구조','Concat 대비 +2.09pp','AD2 16-shot, seed 0'],
 ['다해상도','최선 단일 대비 +4.95 / +2.10pp','AD2 2 / 16-shot, seed 0'],
 ['q95 보정','none 대비 −0.33~+0.40pp','일관된 AUPRO 향상 미확인'],
 ['다음 단계','Local-center → 다중 seed → 외부 평가','최종 결과 미확정']
 ],[2.2,5.4,4.34],y=2.0,row_h=.72,size=16)

s = frame(4,S1,'정상 이미지 수는 그대로, 관측 축을 늘린다','문제 │ 소수 정상 reference로 미세 결함과 정상 변동을 구분해야 한다',
 ['동결 백본으로 정상 패치와 query 패치를 비교한다.', '다해상도·다계층 관측은 같은 이미지의 다른 단서를 제공한다.'])
statement(s,3.1,'공간 관측','448 / 672 / 896 / 1120: 서로 다른 토큰 밀도')
statement(s,4.1,'계층 관측','Layer 5 / 11 / 23: 서로 다른 깊이의 표현')
statement(s,5.1,'해석의 범위','새로운 정상 사례나 센서 정보를 추가하는 것은 아니다.')

s = frame(5,S1,'관측마다 bank와 거리 맵을 유지한다','구조 │ 특징을 먼저 합치는 대신, 각 관측의 점수를 마지막에 결합한다')
rows(s,['단계','특징 결합 대조','현재 방법'],[
 ['관측','여러 계층의 특징','여러 계층·해상도의 특징'],
 ['정상 reference','계층 특징 concat → 하나의 bank','관측 조건별 별도 bank'],
 ['Query 비교','결합 특징의 거리 맵','관측별 거리 맵'],
 ['최종 결합','하나의 거리 점수','Support 단위 보정 후 점수 결합']
 ],[2.1,4.9,4.94],y=2.15,row_h=.80,size=17)
text_box(s,.85,6.0,11.5,.45,'별도 bank 구조의 효과와 q95 연산의 효과는 각각 대조한다.',17,NAVY,BOLD,True)

s = frame(6,S1,'단위 정렬의 동기와 성능 기여를 구분한다','해석 │ 정상 support를 기준으로 거리의 상대 단위를 정의한다',
 ['계층·해상도별 거리 크기는 서로 다를 수 있다.', 'q95는 각 bank의 정상 거리 기준점을 1로 맞추는 연산이다.'])
equation(s,'관측별 거리 dₘ  ÷  정상 기준 cₘ  =  보정 점수 zₘ',3.3,27)
statement(s,4.7,'설계상 역할','서로 다른 관측의 점수를 일정한 규칙으로 결합한다.')
statement(s,5.7,'실험상 상태','none 대비 성능 우위와 query FPR 보장은 확인되지 않았다.')

s = frame(7,S2,'정상 support에서 bank와 척도를 한 번 만든다','피팅 │ 동결 백본의 특징과 정상 거리만 사용한다')
for y,a,b in [(2.15,'1. 특징 추출','4개 해상도 × 3개 계층'),
               (3.12,'2. Bank 구성','해상도·bag별 이미지당 256개 위치, 계층 간 위치 공유'),
               (4.09,'3. 정상 거리','Leave-Source-Out 매칭으로 정상 거리 분포 수집'),
               (5.06,'4. 척도 고정','컴포넌트별 q95: cₘ = Q₀.₉₅(Dₘ)')]:
    statement(s,y,a,b,.78)
text_box(s,.85,6.15,11.5,.4,'Query label을 fitting에 사용하지 않는다. λ 선택은 별도의 개발 평가 기록이다.',15,SUB)

s = frame(8,S2,'Query는 보정된 점수 맵으로 결합한다','추론 │ 특징 추출 → 거리 → 보정 → 계층 결합 → 공간 평균')
rows(s,['단계','연산','출력'],[
 ['1','동결 DINOv2, 4개 입력 해상도','계층별 패치 특징'],
 ['2','각 bank와 1-NN cosine 거리','48개 거리 맵 dₘ'],
 ['3','zₘ = dₘ / cₘ','관측별 보정 맵'],
 ['4','(z₁₁ + z₂₃) / 2 + 0.125 z₅','해상도·bag별 결합 맵'],
 ['5','Bag 평균 → 원본 좌표 보간 → 해상도 평균','최종 맵 · 상위 1% 이미지 점수']
 ],[1.1,6.1,4.74],y=2.0,row_h=.72,size=16)

s = frame(9,S2,'48개 bank가 관측 조건을 분리한다','규격 │ 3개 계층 × 4개 해상도 × 4개 bag',
 ['Cap256은 이미지당·해상도당·bag당 추출하는 토큰 수다.', '16-shot의 한 컴포넌트 bank는 4,096개 토큰으로 구성된다.'])
rows(s,['해상도','짧은 변 토큰 수','정사각 기준 토큰 수'],[
 ['448','32','1,024'],['672','48','2,304'],['896','64','4,096'],['1120','80','6,400']
 ],[3.4,4.2,4.34],y=3.05,row_h=.62,size=17)
text_box(s,.85,6.2,11.5,.38,'각 bank는 같은 support와 백본을 공유하며 통계적 독립성을 가정하지 않는다.',15,SUB)

s = frame(10,S2,'q95는 정상 거리의 단위를 맞춘다','정의 │ LSO 정상 거리의 95% 분위수를 컴포넌트 척도로 사용한다')
equation(s,'cₘ = Q₀.₉₅(Dₘᴸˢᴼ)     ·     zₘ = dₘ / cₘ',2.6,32)
statement(s,4.0,'zₘ = 1','정상 support 기준 거리와 같은 크기')
statement(s,4.9,'zₘ = 2','그 기준 거리의 두 배')
statement(s,5.8,'적용 범위','결함 확률이나 최종 배포 임계값을 뜻하지 않는다.')

s = frame(11,S2,'Shallow를 고정 가중치 0.125로 더한다','결합 │ 계층별 점수 맵을 유지한 뒤 선형 가산한다')
equation(s,'z_d = (z₁₁ + z₂₃) / 2',2.3,31)
equation(s,'S = z_d + 0.125 × z₅',3.3,37)
statement(s,4.65,'가중치 선택','AD2 public 2-shot seed 0 AUPRO로 선택한 post-hoc 값')
statement(s,5.6,'주장 범위','Shallow는 관측 추가 장치이며 독립적인 필수성은 주장하지 않는다.',.8)

s = frame(12,S2,'공간 평균과 상위 1% 점수로 평가한다','출력 │ 해상도별 맵을 원본 크기로 보간한 뒤 동일 가중 평균한다')
rows(s,['순서','처리','역할'],[
 ['1','같은 해상도의 4개 bag 평균','샘플링별 점수를 통합'],
 ['2','각 맵을 원본 H × W로 양선형 보간','공간 좌표를 일치'],
 ['3','4개 해상도 맵의 평균','최종 결함 국소화 맵'],
 ['4','최종 맵의 상위 1% 픽셀 평균','I-AUROC용 이미지 점수']
 ],[1.2,6.1,4.64],y=2.15,row_h=.82,size=17)
text_box(s,.85,6.1,11.5,.4,'비교 실험의 공통 설정: FP16 · cap256 · 4 bags · CQ99 · noise guard k9',15,SUB)

s = frame(13,S3,'AD2 저-FPR 국소화를 주 평가로 둔다','평가 │ 데이터셋별 지표와 seed 범위를 구분해 읽는다')
rows(s,['평가','지표','현재 범위'],[
 ['AD2 public · 8 classes','AUPRO@.05 / pAUROC@.05 / I-AUROC','1 / 2 / 4 / 8 / 16-shot'],
 ['MVTec AD · 15 classes','I-AUROC / P-AUROC 전체 구간','1 / 2 / 4 / 8 / 16-shot'],
 ['VisA · 12 classes','I-AUROC / P-AUROC 전체 구간','1 / 2 / 4 / 8 / 16-shot'],
 ['Ours seed 범위','2·16-shot: 3-seed','1·4·8-shot: seed 0'],
 ['평가 해석','Class macro · %','모두 반복 사용한 개발 모집단']
 ],[3.1,4.5,4.34],y=2.05,row_h=.73,size=15)

au = [r for r in t4 if r[1]=='AUPRO@0.05']
s = frame(14,S3,'AD2 16-shot AUPRO는 60.06 ± 0.88이다','주 결과 │ SuperAD 3-seed 평균 55.89 대비 +4.17pp',
 ['2-shot도 3-seed 평균 56.03 ± 2.99이며, SuperAD 대비 +6.23pp다.'],
 note='AUPRO@0.05 (%) · Ours 2·16-shot만 3-seed, 나머지는 seed 0 · 비교군 3-seed · 표 4',
 refs='표 4. Ours FP16 4-scale, 비교군 FP32. Backbone·관측 수·계산량이 다른 전체 파이프라인 비교다.')
chart(s,shots,[(name,[num(r[col]) for r in au]) for name,col in [('Ours',9),('SuperAD',7),('PatchCore',2),('SubspaceAD',6),('AnomalyDINO',3)]],low=30,high=65,ylabel='AUPRO@0.05 (%)')

s = frame(15,S3,'MVTec AD·VisA의 우위는 지표별로 다르다','범위 │ MVTec AD 고-shot pixel 지표와 VisA 저-shot에서는 비교군이 앞선다',
 ['VisA 16-shot에서는 Ours I/P-AUROC가 95.99 / 98.49로 표의 최고값이다.'],
 note='Seed 0 · I-AUROC / P-AUROC 전체 구간 (%) · 표 5 · 각 방법 설정 상이')
for x,ds in [(.75,'MVTec AD'),(6.9,'VisA')]:
    rr=[r for r in t5 if r[0]==ds]
    text_box(s,x,2.35,5.6,.4,ds + ' · I-AUROC',16,NAVY,BOLD,True)
    chart(s,shots,[(name,[num(r[c]) for r in rr]) for name,c in [('Ours',9),('SuperAD',7)]],
          x=x,w=5.65,y=2.85,h=3.45,low=85,high=100)

s = frame(16,S3,'2·16-shot의 전체 방법은 3-seed를 확보했다','변동 │ AD2 2-shot은 16-shot보다 seed 간 편차가 크다')
rows(s,['데이터셋 · 지표','2-shot','16-shot'],[
 ['AD2 · AUPRO@.05','56.03 ± 2.99','60.06 ± 0.88'],
 ['AD2 · pAUROC@.05','75.86 ± 0.59','77.61 ± 0.37'],
 ['AD2 · I-AUROC','74.93 ± 1.54','75.92 ± 0.82'],
 ['MVTec AD · I / P','96.04 ± 0.30 / 97.38 ± 0.04','98.52 ± 0.20 / 98.05 ± 0.00'],
 ['VisA · I / P','92.38 ± 1.04 / 97.40 ± 0.08','95.78 ± 0.18 / 98.50 ± 0.02']
 ],[3.0,4.47,4.47],y=2.05,row_h=.73,size=15)
text_box(s,.85,6.22,11.5,.4,'Mean ± sample SD (%) · MVTec AD/VisA 추가 seed는 외부 비교군과 support가 짝지어지지 않음',13,SUB)

s = frame(17,S4,'별도 bank는 계층 concat보다 높았다','구조 근거 │ AD2 16-shot AUPRO +2.09pp · VisA I-AUROC +1.90pp',
 ['같은 float 수·cosine 거리의 계층 concat 대조와 비교한 seed 0 결과다.'],
 note='표 8a · AD2 AUPRO@.05, MVTec AD/VisA I-AUROC (%) · seed 0 개발 비교')
rows(s,['데이터셋','Shot','Layer concat','Full','차이 (pp)'],[
 ['AD2','2','52.12','59.33','+7.21'],['AD2','16','57.48','59.57','+2.09'],
 ['MVTec AD','16','98.19','98.74','+0.55'],['VisA','16','94.09','95.99','+1.90']
 ],[2.6,1.2,2.7,2.7,2.74],y=2.8,row_h=.7,size=17)

s = frame(18,S4,'다해상도는 최선 단일보다 AUPRO가 높았다','AD2 │ 2-shot +4.95pp · 16-shot +2.10pp',
 ['각 bank의 cap은 같지만 총 bank 용량은 달라, 용량 효과를 분리한 비교는 아니다.'],
 note='표 10 · linear0125 · q95 · FP16 · seed 0 · AUPRO@0.05 (%)')
chart(s,['448','672','896','1120','4-scale'],[
 ('2-shot',[51.14,54.38,52.65,54.27,59.33]),('16-shot',[52.01,55.01,55.12,57.47,59.57])],low=48,high=62,ylabel='AUPRO@0.05 (%)')

s = frame(19,S4,'다해상도 이득은 MVTec AD·VisA에서도 보인다','해상도 │ 두 데이터셋·두 shot의 I/P-AUROC가 모두 상승했다',
 note='표 10b · I-AUROC / P-AUROC (%) · seed 0 · 총 bank 용량 차이 포함')
rows(s,['Dataset','Shot','1120 단일','4 해상도'],[
 ['MVTec AD','2','94.90 / 96.51','96.13 / 97.39'],
 ['MVTec AD','16','98.21 / 97.49','98.74 / 98.05'],
 ['VisA','2','86.88 / 96.19','92.15 / 97.32'],
 ['VisA','16','94.39 / 97.80','95.99 / 98.49']
 ],[3.0,1.4,3.77,3.77],y=2.2,row_h=.78,size=18)
text_box(s,.85,6.1,11.5,.45,'입력 해상도를 높이기만 하는 경우와 여러 해상도를 결합하는 경우를 구분한다.',16,NAVY,BOLD,True)

s = frame(20,S4,'큰 shallow 가중치는 AD2 성능을 낮춘다','민감도 │ λ=0.125를 개발 평가로 선택하고 다른 데이터셋에도 고정했다',
 ['2-shot seed 0의 AUPRO로 선택한 post-hoc 설정이며, 일반적 최적값으로 주장하지 않는다.'],
 note='표 8b · q95 · FP16 · seed 0 · 양의 가중치 조건 · 후보 전체는 발표자 노트에 기록',
 refs='λ 후보: {0, 0.125, 0.25, 0.5, 1, 2}. 표에는 양의 가중치만 표시. AD2 public 2-shot에서 선택 후 동결.')
chart(s,['0.125','0.25','0.5','1','2'],[
 ('2-shot',[59.33,58.43,54.92,48.39,40.48]),('16-shot',[59.57,59.48,58.21,54.02,47.33])],low=35,high=65,ylabel='AUPRO@0.05 (%)')

s = frame(21,S4,'q95의 none 대비 성능 기여는 확인되지 않았다','보정 │ AD2 AUPRO 차이는 −0.33~+0.40pp로 방향이 엇갈린다',
 ['q95는 MAD보다 높지만, 보정을 제거한 none보다 일관되게 높지는 않다.',
  'MVTec AD·VisA의 I-AUROC도 q95가 none보다 0.07~1.06pp 낮다.'],
 note='표 8d · 전체 방법에서 normalization만 변경 · seed 0 · AUPRO@0.05 (%)')
chart(s,shots,[(name,[num(r[c]) for r in t8d]) for name,c in [('q95',3),('none',1),('MAD',2)]],y=2.95,h=3.45,low=52,high=61,ylabel='AUPRO@0.05 (%)')

s = frame(22,S4,'분위수 변경만으로는 개선이 충분하지 않았다','진단 │ q=0.80~0.999에서 AUPRO 변화는 −0.22~+0.04pp다')
rows(s,['변형','2-shot AUPRO / pAUROC','16-shot AUPRO / pAUROC'],[
 ['q95 기본','59.33 / 76.48','59.57 / 77.46'],
 ['q90','59.29 / 76.42','59.60 / 77.49'],
 ['q99','59.37 / 76.52','59.55 / 77.39'],
 ['꼬리 확률 변환','59.11 / 76.65','59.25 / 77.14'],
 ['토큰별 local scale','52.34 / 75.33','57.46 / 78.03']
 ],[3.3,4.32,4.32],y=2.0,row_h=.63,size=16)
text_box(s,.85,5.9,11.5,.7,'척도 차이는 평균 약 1.8배다. 완만한 상대 가중 조정이라는 해석과 부합하지만,\nnone과 비슷한 성능을 내는 원인으로 확정하지 않는다.',15,SUB)

s = frame(23,S5,'다음 보정은 정상 거리의 국소 중심을 다룬다','후속 │ Local-center는 bank별 단위 보정을 유지하며 거리 수준을 보정한다',
 ['가까운 정상 support의 거리 수준과 전체 평균의 차이를 query 거리에서 뺀다.',
  '강도 0.25 / 0.5를 q95·none과 동일 조건으로 비교한다.'],
 refs='configs/benchmarks/scarf_local_center_fp16_20261012/PROTOCOL.md. Outline 기록 시점 기준 최종 결과 미확정. 라이브 진행률을 반영한 슬라이드가 아니다.')
equation(s,'z = [d − α × (국소 정상 거리 − 전체 평균)] / q95',3.4,26)
statement(s,4.8,'판정 기준','AD2 2·16-shot에서 q95와 none 모두 대비 개선 여부 확인')
statement(s,5.75,'현재 상태','최종 결과 미확정 · 성능 기여로 서술하지 않음')

s = frame(24,S5,'현재 주장할 수 있는 범위를 명확히 한다','논리 │ 관측 확장·bank 구조를 중심으로, 보정 효과는 별도로 검증한다')
rows(s,['주장','근거 상태','발표 표현'],[
 ['전체 방법의 AD2 성능','2·16-shot 3-seed 확보','개발 비교에서 높은 성능'],
 ['관측별 bank 구조','Concat 대조, seed 0','현재 비교 조건에서 지지'],
 ['다해상도 결합','세 데이터셋, seed 0','총 용량 차이를 포함한 이득'],
 ['q95 성능 필수성','none 대비 일관된 이득 없음','기여를 주장하지 않음'],
 ['Shallow 자체의 필수성','현재 원고의 주장 대상 아님','관측 추가의 설계로 설명']
 ],[3.05,4.3,4.59],y=2.05,row_h=.73,size=15)

s = frame(25,S5,'보정 검증 이후 남은 평가를 마무리한다','계획 │ 현재 방법 설정으로 확보하지 않은 결과는 비워 둔다')
rows(s,['우선순위','남은 작업','완료 기준'],[
 ['1','Local-center 보정 결과 취합','q95·none 대조에서 개선 판단'],
 ['2','1·4·8-shot 다중 seed','AD2·MVTec AD·VisA의 통계량 통일'],
 ['3','공식 private · shift · corruption','linear0125 결과와 threshold 검증'],
 ['4','계산 비용 · 구성 요소별 비용','동일 장비·조건에서 절대값과 비율'],
 ['5','비교 감사 · 정성 맵 · 방법 명칭','Support pairing·조건 명시·그림 정비']
 ],[1.45,5.25,5.24],y=2.05,row_h=.74,size=15)

s = frame(26,S5,'성능 근거는 확보했고, 보정 기여가 다음 과제다','결론 │ AD2 중심의 전체 방법과 관측별 bank 구조를 우선 근거로 삼는다')
statement(s,2.2,'전체 성능','AD2 16-shot: 60.06 ± 0.88 AUPRO@0.05 (3-seed)',.8)
statement(s,3.35,'구조 근거','Concat 대비 +2.09pp · 최선 단일 해상도 대비 +2.10pp',.8)
statement(s,4.5,'보정의 한계','q95의 none 대비 일관된 성능 기여는 아직 없다.',.8)
statement(s,5.65,'다음 결정','Local-center를 통제 비교한 뒤 보정 설계를 판단한다.',.8)

# Appendix: transpose tables so mean ± SD and all comparison methods remain legible.
names=['PatchCore','AnomalyDINO','FastRecon','RAD','SubspaceAD','SuperAD','RegAD (224)','Ours']
for n,metric in [(27,'AUPRO@0.05'),(28,'pAUROC@0.05'),(29,'I-AUROC')]:
    rr=[r for r in t4 if r[1]==metric]
    s=frame(n,SA,'AD2 상세 결과 · '+metric,'표 4 │ 비교군 3-seed · Ours 2·16-shot 3-seed, 나머지는 seed 0',
            note='단위 % · mean ± sample SD · RegAD 입력 224 · Ours FP16 4-scale, 나머지 FP32')
    tf=rows(s,['방법']+[x+'-shot' for x in shots],[[name]+[r[c] for r in rr] for c,name in enumerate(names,2)],
         [2.25]+[1.938]*5,y=1.95,row_h=.49,size=12)
    rank_columns(tf)
    s.notes_slide.notes_text_frame.text += '\n각 shot 최고 평균은 굵게, 차상위 평균은 밑줄. 1/4/8-shot Ours는 단일 seed이므로 순위는 동일 통계량 비교가 아니다.'

s=frame(30,SA,'AD2 비교군은 최대 입력을 맞춘 설정이다','표 3 │ Backbone·관측 수·정밀도는 달라 전체 파이프라인 비교로 해석한다')
rows(s,['방법','Backbone · layers','입력 / 설정'],[
 ['Ours','DINOv2 ViT-L/14 · 5/11/23','448/672/896/1120 · FP16'],
 ['PatchCore','DINOv2 ViT-L/14 · 5/11/17/23','1120 · FP32'],
 ['SuperAD','DINOv2 ViT-L/14 · 5/11/17/23','1120 · FP32'],
 ['AnomalyDINO','Project DINOv2 ViT-L/14 · 23','1120 · FP32'],
 ['FastRecon','ResNet-50 · paper configuration','1120 · FP32 · 独立 재구현'.replace('独立','독립')],
 ['RAD','DINOv3 ViT-B/16 · 3/6/9/11','1280 → crop1120 · FP32'],
 ['SubspaceAD','DINOv2-reg ViT-G/14 · 22–28','1120 · FP32'],
 ['RegAD','ResNet-18 STN · source checkpoint','224 · 공식 고정 입력']
 ],[2.05,5.5,4.39],y=1.95,row_h=.5,size=13)
s.notes_slide.notes_text_frame.text += '\nFastRecon/RAD 독립 재구현. SuperAD는 원 논문 reference 선택과 달리 공통 support 사용. 표 5의 E9 FastRecon DINOv2 변형과 표 4의 ResNet-50 설정을 구분한다.'

for n,ds in [(31,'MVTec AD'),(32,'VisA')]:
    rr=[r for r in t5 if r[0]==ds]
    s=frame(n,SA,ds+' 상세 결과 · I / P-AUROC','표 5 │ 모든 셀 seed 0 단일값 · P-AUROC는 전체 FPR 구간',
        note='단위 % · E9 기존 설정과 신규 1120 설정 혼재 · linear0125 support hash 재감사 필요')
    tf=rows(s,['방법']+[x+'-shot' for x in shots],[[name]+[r[c] for r in rr] for c,name in enumerate(names,2)],
         [2.25]+[1.938]*5,y=1.95,row_h=.49,size=12)
    rank_columns(tf,pairs=True)
    s.notes_slide.notes_text_frame.text += '\n표 5: PatchCore/SuperAD/FastRecon은 E9 설정. 신규 AnomalyDINO/RAD/SubspaceAD는1120. FastRecon은 DINOv2 프로젝트 변형. 모든 비교가 동일 입력이라는 주장은 하지 않는다.'

s=frame(33,SA,'AD2 보정 factorial · 전체 방법','표 8d │ Normalization만 변경 · 각 셀 AUPRO / pAUROC / I-AUROC',
        note='단위 % · AD2 seed 0 full8 · q95의 none 대비 일관된 성능 이득 없음')
rows(s,['Shot','none','MAD','q95'],[[r[0],r[1],r[2],r[3]] for r in t8d],
     [1.0,3.64,3.64,3.66],y=2.15,row_h=.65,size=14)
text_box(s,.85,6.15,11.5,.4,'AD2 q95 − none AUPRO: −0.25 / +0.27 / −0.33 / −0.01 / +0.40pp',15,NAVY,BOLD,True)

s=frame(34,SA,'분위수와 보정 변형 상세 결과','표 11 │ AD2 2·16-shot, seed 0 · 각 셀 AUPRO / pAUROC')
variant_names=['q = 0.50','q = 0.80','q = 0.90','q = 0.95 (기본)','q = 0.99','q = 0.999','꼬리 확률 변환','토큰별 local scale']
rows(s,['척도 / 변형','2-shot','16-shot'],[[variant_names[i],r[1],r[2]] for i,r in enumerate(t11)],
     [4.0,3.97,3.97],y=1.95,row_h=.5,size=14)

s=frame(35,SA,'재현 설정과 결과의 경계를 고정한다','규격 │ 같은 설정의 linear0125 결과만 현재 방법의 근거로 사용한다')
rows(s,['항목','고정 설정 / 기록'],[
 ['Backbone / precision','Frozen DINOv2 ViT-L/14 · FP16 autocast'],
 ['관측 / bank','Layers 5/11/23 · 448/672/896/1120 · cap256 · 4 bags'],
 ['거리 / 보정','L2-normalized features · 1-NN cosine · support q95'],
 ['결합','(z11 + z23)/2 + 0.125 × z5 · bag/scale 평균'],
 ['평가 옵션','CQ99 · noise guard k9 · image score = top-1% mean'],
 ['선택 기록','λ: AD2 public 2-shot seed 0 post-hoc 선택'],
 ['미확보 결과','현재 방법의 private · shift · corruption · 계산 비용'],
 ['원본 추적','Outline_v17.1.md 표 4/5/5b/8a/8d/10/10b/11']
 ],[2.6,9.34],y=1.95,row_h=.5,size=14)

set_sections(prs,[('표지와 요약',list(prs.slides)[:3]),(S1,list(prs.slides)[3:6]),
                 (S2,list(prs.slides)[6:12]),(S3,list(prs.slides)[12:16]),
                 (S4,list(prs.slides)[16:22]),(S5,list(prs.slides)[22:26]),(SA,list(prs.slides)[26:])])
fix_table_ids(prs)
prs.core_properties.title='SCARF · Outline v17.1 · linear0125'
prs.core_properties.subject='관측별 bank, 다해상도, support score calibration'
prs.core_properties.comments='Updated from Outline_v17.1.md on 2026-10-12. Original v16.3 retained.'
scratch=BUILD / 'SCARF_v17.1_build.pptx'
prs.save(scratch)
shutil.copy2(scratch,OUT)
manifest={
 'source_deck':str(SRC.relative_to(ROOT)), 'source_deck_sha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),
 'outline':str(MD.relative_to(ROOT)), 'outline_sha256':hashlib.sha256(MD.read_bytes()).hexdigest(),
 'output':str(OUT.relative_to(ROOT)), 'slides':len(prs.slides),
 'scope':'v17.1 linear0125 only; no layer-removal comparison; prior gated results excluded',
 'tables_parsed':['4','5','8a','8d','11'],
 'notes':'Source reports retained; pending local-center described as unconfirmed, no live progress claimed.'
}
(BUILD / 'provenance.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print(OUT)
