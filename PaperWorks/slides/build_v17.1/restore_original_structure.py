"""Keep the 35-slide v16.3 topic order and body roles; refresh v17.1 evidence."""
from pathlib import Path
import io
from datetime import datetime
helper = Path(__file__).with_name('build_deck.py')
source = helper.read_text()
exec(compile(source.split('# 1. Cover')[0].replace('else XL_CHART_TYPE.COLUMN_CLUSTERED', "else (XL_CHART_TYPE.COLUMN_STACKED if kind == 'stack' else XL_CHART_TYPE.COLUMN_CLUSTERED)"), str(helper), 'exec'))
from deck_ops import set_lines, shape
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_DATA_LABEL_POSITION
from copy import deepcopy

stamp=datetime.now().strftime('%Y%m%d_%H%M%S')
bak=OUT.parent/'backup'/(OUT.stem+'_before_structure_restore_'+stamp+'.pptx')
shutil.copy2(OUT,bak)
if OUT.with_suffix('.pdf').exists(): shutil.copy2(OUT.with_suffix('.pdf'),bak.with_suffix('.pdf'))
(BUILD/'structure_backup_path.txt').write_text(str(bak))
original = Presentation(SRC)
S1,S2,S3,S4,S5='01. 문제 정의','02. 방법: SCARF','03. 실험 결과','04. Ablation · 분석','05. 논의와 향후 계획'
SA='Appendix A. 상세 표'

def frame(n,sec,title,lead,detail=(),note='',refs=''):
    s=prs.slides[n-1]
    keep={'SectionLabel','Title','Logo','Lead','MetricNote','PageNumber'}
    for sh in list(s.shapes):
        if sh.name not in keep: sh._element.getparent().remove(sh._element)
    for nm,txt in [('SectionLabel',sec),('Title',title),('Lead',lead),('PageNumber',str(n) if n<=26 else 'A-'+str(n-26))]:
        set_lines(shape(s,nm).text_frame,txt)
    notes=[x for x in s.shapes if x.name=='MetricNote']
    if notes: set_lines(notes[0].text_frame,note or source_footer)
    else: text_box(s,.7,6.83,11.94,.3,note or source_footer,9,SUB,name='MetricNote')
    for i,txt in enumerate(detail):
        text_box(s,.8,1.78+i*.39,11.7,.37,'• '+txt,16,INK,name='Detail')
    s.notes_slide.notes_text_frame.text=(f'출처: PaperWorks/outlines/Outline_v17.1.md\n{title}\n{lead}\n'+'\n'.join(detail)+'\n'+(note or source_footer)+'\n'+refs+'\n개발에 반복 사용한 평가 모집단이며 untouched holdout 주장이 아니다. 이전 gated SCARF 결과를 현재 linear0125 결과로 재사용하지 않는다.')
    return s

def reuse(old,new,end):
    start=re.search(r'^s ?= ?frame\('+str(old)+',',source,re.M).start()
    stop=source.index(end,start)
    code=source[start:stop].replace('frame('+str(old)+',','frame('+str(new)+',',1)
    exec(code,globals())

def picture_from(n,s,which=0):
    pics=[p for p in original.slides[n-1].shapes if p.name=='ContentFigure']
    p=pics[which]
    return s.shapes.add_picture(io.BytesIO(p.image.blob),p.left,p.top,p.width,p.height)

def label(s,x,y,w,txt,size=17):
    return text_box(s,x,y,w,.5,txt,size,NAVY,BOLD,True,align=PP_ALIGN.CENTER)

def box(s,x,y,w,h,txt,size=18):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.background(); sh.line.color.rgb=RGBColor.from_string('ADB8C6');sh.line.width=Pt(.8)
    text_box(s,x+.1,y+.16,w-.2,h-.2,txt,size,INK,align=PP_ALIGN.CENTER)
    return sh

def arrow(s,x,y): label(s,x,y,.35,'→',21)

def pending(n,title,lead,items,labels,refs):
    s=frame(n,S3 if n<19 else S4,title,lead,items,note='linear0125 현재 결과 미확보 · 기존 구성과 평가 항목 유지 · '+refs)
    width=(11.5-.22*(len(labels)-1))/len(labels)
    for j,t in enumerate(labels):
        x=.9+j*(width+.22)
        label(s,x,3.05,width,t,17)
        box(s,x,3.65,width,1.6,'—\n현재 설정 결과 미확보',17)
    text_box(s,.9,5.8,11.5,.6,'현재 방법의 실측 결과 확보 후 동일 위치에 차트를 반영한다.',17,SUB)
    return s

# Cover: preserve every original position.
s=prs.slides[0]
for name,txt in [('Subtitle','관측별 Memory Bank와 Support 기반 점수 정렬'),('Subtitle2','소수 정상 이미지 기반 산업 결함 국소화 · linear0125 설정'),('Date','연구 아웃라인 v17.1 발표  ·  2026-10-12')]: set_lines(shape(s,name).text_frame,txt)
s.notes_slide.notes_text_frame.text='Outline_v17.1.md 기준. SCARF의 기존 Asymmetric Residual Fusion 확장명은 변경된 선형 결합에 맞지 않아 재검토 중. v16.3의 표지·35장 구성 유지.'

s=frame(2,'목차','발표 순서','흐름 │ 문제 → 방법 → 주 결과 → 소거·분석 → 남은 검증 순으로 진행한다')
for i,t in enumerate(['01  문제 정의: 소수 정상 관측과 점수 단위의 차이 (4–7쪽)','02  방법 SCARF: 별도 bank·척도 보정·선형 결합 (8–12쪽)','03  주 결과: AD2·MVTec AD·VisA와 미확보 평가 (13–18쪽)','04  소거·분석: 구조·보정·해상도·계산 비용 (19–23쪽)','05  논의와 향후 계획: 근거의 범위와 남은 검증 (24–26쪽)','Appendix A  상세 수치와 고정 설정 (A-1–A-9)']):
    text_box(s,.85,1.85+i*.68,11.5,.5,'• '+t,19,INK)

s=frame(3,'요약','한 장 요약','결론 │ 관측별 bank와 다해상도 결합의 근거는 확보했고, 보정 기여는 추가 검증한다')
rows(s,['섹션','질문','결론','쪽'],[
 ['문제','소수 정상 관측을 어떻게 보완하나?','같은 이미지의 해상도·계층 관측 확장','4–7'],
 ['방법','서로 다른 점수를 어떻게 합치나?','별도 bank → q95 → +0.125 shallow','8–12'],
 ['주 결과','AD2 성능은 어느 수준인가?','16-shot AUPRO 60.06 ± 0.88 (3-seed)','13–18'],
 ['소거·분석','어떤 구조가 성능을 뒷받침하나?','Concat·단일 해상도 대비 이득; q95는 미확정','19–23'],
 ['다음 검증','무엇을 보완해야 하나?','보정 개선·다중 seed·외부 평가·비용','24–26']
 ],[1.4,3.0,6.44,1.1],y=1.8,row_h=.73,size=15)

s=frame(4,S1,'배경: 정상 몇 장으로 결함 위치 찾기','목적 │ 정상 1–16장으로, 오경보가 적은 구간에서 결함 영역을 찾는다',[
 '동결 DINOv2 패치를 정상 bank와 비교하고, 가장 가까운 정상 패치와의 거리를 계산한다.',
 '다해상도·다계층은 같은 이미지의 관측 축을 늘리며, 새로운 정상 사례를 추가하지는 않는다.'])
for x,t,b in [(1.,'정상 support','정상 1–16장\n해상도·계층별 bank'),(5.,'검사 이미지','동결 특징 추출\n관측별 거리 맵'),(9.,'결함 국소화','보정·선형 결합\n최종 anomaly map')]:
    label(s,x,3.1,3.2,t,20);box(s,x,3.75,3.2,1.4,b,20)
arrow(s,4.45,4.18);arrow(s,8.45,4.18)
label(s,1.,5.75,11.2,'주 지표: AUPRO@0.05 — 낮은 FPR 구간의 결함 영역 국소화',19)

s=frame(5,S1,'문제 ①: 큰 거리가 항상 결함은 아니다','원인 │ 소수 정상 support는 조명·자세·반사 등 정상 변동을 모두 담지 못한다',[
 '처음 보는 정상 조건과 결함이 모두 높은 거리 점수를 만들 수 있다.',
 '관측별 점수를 결합하려면, 각 bank에서 거리 크기가 갖는 의미부터 정렬해야 한다.'],note='가상 분포 개념도이며 실측 결과가 아님 · v16.3 개념 그림 유지 · Outline_v17.1 §1')
picture_from(5,s)

reuse(20,6,'s = frame(21,')
set_lines(shape(prs.slides[5],'SectionLabel').text_frame,S1)
set_lines(shape(prs.slides[5],'Title').text_frame,'문제 ②: 얕은 층의 기여량을 조절해야 한다')

s=frame(7,S1,'해법의 틀: 단위와 기여량을 따로 정한다','핵심 │ 관측별 점수를 정상 support 척도로 정렬하고, shallow 기여는 고정 가중치로 더한다',[
 '척도 보정: 각 bank의 정상 거리 q95를 기준으로 서로 다른 거리 눈금을 맞춘다.',
 '기여량: shallow 가중치 0.125는 AD2 개발 평가로 선택하고 다른 데이터셋에도 고정한다.'])
label(s,1.,3.1,5.3,'① 관측별 단위 정렬',21);label(s,7.,3.1,5.3,'② 계층별 점수 결합',21)
text_box(s,1.,3.9,5.3,.75,'zₘ = dₘ / Q₀.₉₅(Dₘ)',27,NAVY,BOLD,True,align=PP_ALIGN.CENTER)
text_box(s,6.9,3.9,5.4,.75,'S = z_d + 0.125 z₅',27,NAVY,BOLD,True,align=PP_ALIGN.CENTER)
text_box(s,1.,5.3,11.4,.75,'q95의 단위 정렬 역할과 실제 성능 기여는 구분한다.\nnone 대비 일관된 성능 향상은 아직 확인되지 않았다.',18,INK,align=PP_ALIGN.CENTER)

s=frame(8,S2,'SCARF 전체 흐름','구조 │ 정상 support로 bank와 기준값을 한 번 만들고(A), query마다 같은 규칙으로 결합한다',[
 'A 피팅: 계층 3 × 해상도 4 × bag 4 = 48개 bank와 정상 거리 척도를 구성한다.',
 'B 추론: 거리 → q95 보정 → 선형 결합 → bag·해상도 평균 순으로 최종 맵을 만든다.'])
for y,head,parts in [(2.85,'A  피팅 (1회)',['정상 support\n동결 특징','관측별 bank\ncap256 × shot','정상끼리 매칭\nLSO 거리','척도 고정\ncₘ = q95']), (4.8,'B  Query 추론',['4 해상도\n3 계층 특징','bank별 거리\nzₘ = dₘ / cₘ','선형 결합\nz_d + 0.125 z₅','bag·scale 평균\n상위 1% 점수'])]:
    text_box(s,.8,y,11.7,.4,head,18,NAVY,BOLD,True)
    for j,t in enumerate(parts):
        x=.8+j*3.03;box(s,x,y+.48,2.65,1.04,t,17)
        if j<3: arrow(s,x+2.68,y+.78)

s=frame(9,S2,'관측 구성: 4개 해상도 × 4개 bag','설계 │ 해상도·계층별 bank를 유지하고, 같은 해상도의 bag 점수는 평균한다',[
 '해상도: 448 / 672 / 896 / 1120을 모두 사용해 공간 관측을 확장한다.',
 'Bag: 이미지당 256개 위치를 4번 추출하며, 계층 간에는 같은 위치를 공유한다.'],note='관측 구성 도식 · 실제 anomaly map 아님 · FP16 · Layer 5/11/23 · Outline_v17.1 §3')
for j,(r,grid) in enumerate([(448,32),(672,48),(896,64),(1120,80)]):
    x=.85+j*3.0;label(s,x,2.75,2.65,str(r)+' px',20)
    # An illustrative patch lattice, not a measured map.
    for k in range(9):
        xx=x+.62+k*.17
        ln=s.shapes.add_connector(1, Inches(xx), Inches(3.3), Inches(xx), Inches(4.35));ln.line.color.rgb=RGBColor.from_string('ADB8C6')
    for k in range(7):
        yy=3.3+k*.175
        ln=s.shapes.add_connector(1, Inches(x+.62), Inches(yy), Inches(x+1.98), Inches(yy));ln.line.color.rgb=RGBColor.from_string('ADB8C6')
    label(s,x,4.42,2.65,'짧은 변 '+str(grid)+' tokens',15)
for j in range(4): box(s,.85+j*2.3,5.25,1.9,.7,'Bag '+str(j+1),17)
arrow(s,10.0,5.43);label(s,10.45,5.38,1.9,'평균 맵',19)

s=frame(10,S2,'Scale Calibration: 거리 단위 통일','방법 │ 각 거리를 정상 support 거리의 95% 분위수 cₘ으로 나눈다',[
 '같은 거리도 bank별 정상 분포가 다르면 다른 크기의 이탈을 뜻한다.',
 'z = 1은 support 척도의 기준점이며, query의 FPR이나 결함 확률을 보장하지 않는다.'],note='기존 pre-fusion 척도 설명 그림: hazelnut 1120 px · 현재 전체 방법의 성능 증거가 아님 · §3.3')
picture_from(10,s)

s=frame(11,S2,'Linear Fusion: Shallow를 0.125만큼 더한다','방법 │ S = z_d + 0.125 z₅,  z_d = (z₁₁ + z₂₃) / 2',[
 '약한 양의 가중치로 shallow 관측을 추가하며, 정상 질감 반응도 함께 더해진다.',
 '0.125는 AD2 public 2-shot seed 0 개발 평가에서 선택한 post-hoc 값이다.'],note='설명용 예시값, 실측 아님 · Outline_v17.1 §3.4 · 임계값 기반 절단 없음')
ch=chart(s,['정상 평탄','정상 질감','미세 스크래치','큰 파손'],[('z_d',[.55,.65,.85,2.4]),('0.125 z₅',[.05,.15625,.35,.375])],y=2.9,h=3.6,low=0,high=3,kind='stack',ylabel='설명용 점수')


s=frame(12,S2,'실제 맵에서 본 선형 결합: 후속 확인','상태 │ 현재 설정의 검증된 정성 맵은 미확보이며, 같은 위치에 비교 패널을 유지한다',[
 '동일 query·동일 색 범위에서 shallow 항이 최종 맵에 추가하는 위치를 확인한다.',
 '정성 패널은 linear0125·4-scale·cap256 조건의 결과로 갱신한다.'],note='정성 맵 미확보 · 아래는 패널 배치 안내이며 anomaly map이 아님 · Outline_v17.1 §3–4')
for j,t in enumerate(['Query / GT','z_d','z₅','0.125 z₅','최종 S']):
    x=.75+j*2.43;label(s,x,2.95,2.2,t,18);box(s,x,3.55,2.2,1.6,'미확보',17)
text_box(s,.9,5.6,11.5,.8,'최종 출력: bag 평균 → 원본 좌표 보간 → 해상도 평균\n이미지 점수: 최종 맵 상위 1% 픽셀 평균',18,INK,align=PP_ALIGN.CENTER)

reuse(13,13,'au =')
au=[r for r in t4 if r[1]=='AUPRO@0.05']
reuse(14,14,'s = frame(15,')
s=frame(15,S3,'MVTec AD · VisA few-shot','범위 │ 우위는 지표·shot별로 다르며, AD2 저-FPR 국소화가 주 평가다',[
 'VisA 16-shot Ours I/P-AUROC는 95.99 / 98.49로 표의 최고값이다.',
 'MVTec AD 고-shot pixel 지표와 VisA 저-shot 일부에서는 비교군이 앞선다.'],note='표 5 · seed 0 · P-AUROC 전체 FPR · 비교군 설정 혼재 · support hash 재감사 필요')
for j,(ds,k,lo) in enumerate([('MVTec AD',0,85),('MVTec AD',1,95),('VisA',0,80),('VisA',1,94)]):
    rr=[r for r in t5 if r[0]==ds];x=.72+j*3.01
    label(s,x,2.7,2.93,ds+' · '+('I' if k==0 else 'P'),15)
    chart(s,shots,[(nm,[num(r[c].split('/')[k]) for r in rr]) for nm,c in [('Ours',9),('SuperAD',7),('PatchCore',2)]],x=x,w=2.93,y=3.1,h=3.48,low=lo,high=100)

pending(16,'MVTec AD 2 공식 private 평가','상태 │ linear0125의 공식 private 결과와 threshold 평가는 아직 없다',[
 '공식 평가의 AU-PRO·ClassF1·SegF1을 함께 보고할 예정이다.',
 'Full-data 참고값은 데이터 예산이 달라 few-shot 방법과 직접 순위화하지 않는다.'],['AU-PRO@0.05','ClassF1','SegF1'],'표 6 · 상세 A-5')
pending(17,'조명 변화 (AD2 Shift 1–3)','상태 │ 현재 방법의 조명별 분해 결과는 미확보다',[
 '정상 regular support를 기준으로 Regular와 Shift 1–3을 구분한다.',
 '조명별 클래스 구성과 macro 범위를 명시한 후 동일 shot 곡선으로 비교한다.'],['Regular','Shift 1','Shift 2','Shift 3'],'§4.3')
pending(18,'합성 손상 (MVTec AD corruption)','상태 │ 현재 방법의 corruption 결과는 미확보다',[
 '16-shot·seed 0·severity 3, 정상 support는 clean 조건을 유지한다.',
 'Clean 대비 저하량을 비교해 밝기·대비·흐림·노이즈 강건성을 평가한다.'],['Clean','Brightness','Contrast','Defocus','Gaussian'],'표 7 · 상세 A-6')

s=frame(19,S4,'구성 요소 소거 (AD2 16-shot)','결과 │ 계층 concat 대비 +2.09pp이며, q95의 none 대비 차이는 +0.40pp다',[
 'Full과 concat은 같은 nominal float 수·cosine 거리의 bank 구조 대조다.',
 '보정의 기여는 모든 shot에서 일관되지 않으므로 이 한 조건으로 일반화하지 않는다.'],note='표 8 · AD2 test_public full8 · seed 0 · AUPRO@0.05 (%) · 상세 A-7')
chart(s,['Full','none','Equal mean','Concat','MAD','448'],[('AUPRO',[59.57,59.17,58.21,57.48,58.52,52.01])],kind='bar',y=2.85,h=3.6,low=48,high=62)

s=frame(20,S4,'저샷·다른 데이터셋 진단','관찰 │ 별도 bank는 concat보다 높지만, q95 보정은 none보다 일관되게 높지 않다',[
 '세 데이터셋의 전 shot에서 Full이 layer concat보다 높다.',
 'MVTec AD·VisA I-AUROC에서는 none이 Full보다 높아 구조와 보정의 효과를 분리한다.'],note='표 8a · seed 0 · AD2 AUPRO@.05, MVTec AD/VisA I-AUROC (%) · 상세 A-7')
for j,(ds,lo) in enumerate([('AD2',40),('MVTec AD',85),('VisA',80)]):
    rr=[r for r in t8a if r[0]==ds];x=.75+j*4.0
    label(s,x,2.7,3.85,ds,17)
    chart(s,shots,[(nm,[num(r[c]) for r in rr]) for nm,c in [('Full',2),('none',3),('Concat',5)]],x=x,w=3.85,y=3.12,h=3.45,low=lo,high=65 if j==0 else 100)
reuse(21,21,'s = frame(22,')
reuse(18,22,'s = frame(19,')
set_lines(shape(prs.slides[21],'Title').text_frame,'해상도 구성 비교')
pending(23,'계산 비용','상태 │ linear0125의 전체 비용과 구성 요소별 비용은 아직 계측하지 않았다',[
 'Fit·query 시간과 VRAM을 같은 장비·정밀도·입력 조건에서 계측해야 한다.',
 '특징 추출·거리 탐색·보정·결합의 절대 시간과 비율을 함께 보고한다.'],['전체 fit / query','Backbone / 탐색','보정 / 결합'],'표 12·12a · 상세 A-8')

s=frame(24,S5,'논의: 관찰과 시사점','시사점 │ 전체 성능·bank 구조·관측 확장은 근거가 있으나, 보정의 기여는 아직 남아 있다',[
 '성능: AD2 16-shot 3-seed AUPRO 60.06 ± 0.88 → 주 평가의 경쟁력 확보.',
 '구조: concat 대조와 단일 해상도 대조의 이득 → 관측별 bank 유지의 근거.',
 '보정: none 대비 일관된 이득 없음 → 국소 정상 거리 수준을 다루는 후속 비교.',
 '검증: 나머지 shot 다중 seed·private·shift·corruption·비용은 추가 확보 필요.'],note='Outline_v17.1 §5 · 모든 근거는 반복 사용한 개발 평가 · 다해상도 총 용량 차이 포함')
box(s,.7,4.25,11.94,1.23,'주장 범위  같은 이미지의 해상도·계층 관측을 별도 bank로 유지하고 점수 단계에서 결합한다.\nq95가 성능에 필수라는 주장은 현재 근거로 확정하지 않는다.',19)
reuse(25,25,'s = frame(26,')
s=frame(26,S5,'결론 및 향후 진행 방향','결론 │ AD2 중심 성능과 관측별 bank 구조를 근거로 삼고, 보정 기여를 다음으로 검증한다')
rows(s,['구분','Summary','현재 문제','향후 진행 방향'],[
 ['방법','별도 bank + 선형 결합','보정의 일관된 이득 미확인','Local-center vs q95 / none'],
 ['성능','AD2 16-shot 60.06 ± 0.88','1·4·8-shot은 단일 seed','전체 shot 다중 seed'],
 ['일반화','MVTec AD·VisA 결과 확보','Private·shift·corruption 미확보','현재 설정의 외부 평가'],
 ['효율','FP16·4-scale·cap256 고정','현재 방법의 실측 비용 미확보','전체·구성 요소 비용 계측']
 ],[1.2,3.9,3.42,3.42],y=1.8,row_h=.9,size=15)

# Appendix preserves original row/column orientation and topic assignment.
names=['PatchCore','AnomalyDINO','FastRecon','RAD','SubspaceAD','SuperAD','RegAD\n(224)','Ours']
widths=[.58]+[1.42]*8

def ad2_table(s,metric,y,h):
    rr=[r for r in t4 if r[1]==metric]
    vals=[]
    for r in rr:
        vals.append([shots[len(vals)]]+[v.replace(' ± ','±' if h<.5 else '\n± ') for v in r[2:]])
    return rows(s,['Shot']+[n.replace('\n',' ') for n in names],vals,widths,y=y,row_h=h,size=10.5)

s=frame(27,SA,'주 비교 원본: AUPRO@0.05','원본 │ 14쪽 차트의 수치 (AD2 test_public, 8-class macro, %)',note='표 4 · 비교군 3-seed · Ours 2·16-shot 3-seed, 1·4·8-shot seed 0 · mean ± sample SD')
ad2_table(s,'AUPRO@0.05',1.8,.68)
s=frame(28,SA,'보조 지표: pAUROC@0.05 · I-AUROC','원본 │ AD2 보조 지표도 같은 seed 범위와 방법 조건으로 해석한다',note='표 4 · 비교군 3-seed · Ours 2·16-shot 3-seed, 나머지 seed 0 · 단위 %')
for y,metric in [(1.65,'pAUROC@0.05'),(4.15,'I-AUROC')]:
    text_box(s,.7,y,6.,.3,metric,15,NAVY,BOLD,True)
    ad2_table(s,metric,y+.32,.35)
reuse(30,29,'for n,ds in')
s=frame(30,SA,'MVTec AD · VisA 원본','원본 │ 15쪽 차트의 수치 (각 셀 I-AUROC / P-AUROC, 전체 FPR, %)',note='표 5 발췌 · seed 0 · 비교군 설정 혼재 · FastRecon DINOv2 프로젝트 변형 · support hash 재감사 필요')
rows(s,['Dataset','Shot','PatchCore','FastRecon','SuperAD','Ours'],[[r[0],r[1],r[2],r[4],r[7],r[9]] for r in t5],[1.45,.65,2.46,2.46,2.46,2.46],y=1.75,row_h=.45,size=12)

s=frame(31,SA,'공식 private 평가 전체 표','참고 │ Full-data 참고값과 few-shot 현재 방법은 데이터 예산이 달라 직접 순위화하지 않는다',note='표 6 · 각 셀 TESTpriv / TESTpriv,mix (%) · current linear0125 결과 미확보')
h6,t6=mtable('6')
# Keep all source methods, combine absent few-shot comparator rows for readable reference table.
data=[['Few-shot: PatchCore / AnomalyDINO / SuperAD / RAD / SubspaceAD','—','—','—'],['Ours (linear0125), 16-shot','— / —','— / —','— / —']]+[['Full-data · '+r[0]]+r[1:] for r in t6[8:]]
rows(s,['방법 / 데이터 예산']+h6[1:],data,[4.8,2.38,2.38,2.38],y=1.75,row_h=.37,size=11)
s.notes_slide.notes_text_frame.text+='\nFull-data 참조값은 원고 표 6의 출판/리더보드 값 인용. 이번 편집에서 외부 최신 순위를 재조사한 것은 아니다.'

s=frame(32,SA,'합성 손상 원본','원본 │ 각 셀 AUPRO@0.05 / pAUROC@0.05 (%) · 현재 방법은 미확보',note='표 7 · MVTec AD 15 classes · 16-shot seed 0 · severity 3 · 비교군 역사적 설정')
h7,t7=mtable('7')
sel=[t7[0],t7[2],t7[5],t7[7]]
sel[1][0]='FastRecon†';sel[3][0]='Ours (linear0125)'
rows(s,['방법','Clean','Brightness','Contrast','Defocus','Gaussian'],sel,[2.14]+[1.96]*5,y=1.8,row_h=.7,size=12)
text_box(s,.85,5.7,11.5,.65,'† DINOv2 기반 프로젝트 변형. 나머지 비교군의 미확보 셀은 원고 표 7에 유지한다.',15,SUB)

s=frame(33,SA,'소거·진단 원본','원본 │ 19·20쪽의 구성 소거와 보정 대조 (seed 0, %)',note='표 8·8a · 아래: AD2 AUPRO/pAUROC, MVTec AD·VisA I/P-AUROC · 2·16-shot, seed 0')
rows(s,['AD2 16-shot','Full','none','Equal mean','Concat','MAD','448'],[
 ['AUPRO','59.57','59.17','58.21','57.48','58.52','52.01'],
 ['pAUROC','77.46','77.65','76.31','74.80','77.17','76.52']], [2.1]+[1.64]*6,y=1.8,row_h=.55,size=12)
rows(s,['Dataset','Shot','Full','none','Concat'],[[r[0],r[1],r[2],r[3],r[5]] for r in t8a if r[1] in ['2','16']],[1.65,.85,3.146,3.147,3.147],y=3.85,row_h=.39,size=12)

s=frame(34,SA,'비교군 계산 비용 (A6000)','상태 │ 현재 방법은 미계측이며, SuperAD 기존 비용은 참고값으로 유지한다',note='표 12·12a · SuperAD 역사적 FP32 조건 · Ours FP16 비용 및 정밀도 일치 비교 필요')
rows(s,['방법','Bank vectors','Fit (s)','Query (ms)','VRAM fit / query'],[
 ['SCARF linear0125','196,608','—','—','—'],['1120-only linear0125','49,152','—','—','—'],
 ['SuperAD (기존 FP32)','All tokens','10.48 ± 2.17','793.95 ± 154.16','1.42 / 1.42 GiB']],[3.1,2.0,2.1,2.54,2.2],y=1.8,row_h=.75,size=12)
text_box(s,.85,5.45,11.5,.8,'추가 계측: cap256 / cap1024 paired 비용, 전체 시간과 구성 요소별 절대값·비율.\n과거 fusion 설정의 시간을 현재 linear0125 비용으로 치환하지 않는다.',16,INK)
reuse(35,35,'set_sections(prs,')


# Retain original topical titles and carry additional v17.1 details in notes.
for n,title in [(13,'평가 설정'),(14,'Few-shot 주 비교 (AD2 test_public)'),(21,'Score Calibration 동작·성능 진단'),(25,'남은 검증'),(29,'AD2 비교군 설정'),(35,'SCARF 고정 사양')]:
    set_lines(shape(prs.slides[n-1],'Title').text_frame,title)
prs.slides[14].notes_slide.notes_text_frame.text += '\n표 5b, Ours 3-seed: MVTec AD 2-shot 96.04±0.30 /97.38±0.04, 16-shot98.52±0.20 /98.05±0.00. VisA 2-shot92.38±1.04 /97.40±0.08,16-shot95.78±0.18 /98.50±0.02. 추가 seed는 외부 비교군과 paired support 아님.'
prs.slides[20].notes_slide.notes_text_frame.text += '\n표 11 보정 변형(2-shot;16-shot AUPRO/pAUROC):\n'+'\n'.join(' | '.join(r[:3]) for r in t11)+'\n컴포넌트 q95 최대/최소 평균 약1.8배. Local-center 최종 결과는 원고 기준 미확정.'

# Row-wise best/runner-up emphasis for appendix comparisons, per metric.
for n,first,pair in [(27,1,False),(28,1,False),(30,2,True)]:
    for sh in prs.slides[n-1].shapes:
        if not sh.has_table: continue
        tb=sh.table
        for i in range(1,len(tb.rows)):
            strings=[tb.cell(i,j).text for j in range(first,len(tb.columns))]
            parts=[t.split('/') if pair else [t] for t in strings]
            ranks=[sorted(set(num(v[k]) for v in parts),reverse=True)[:2] for k in range(2 if pair else 1)]
            for j,pp in enumerate(parts,first):
                tf=tb.cell(i,j).text_frame
                sz=tf.paragraphs[0].runs[0].font.size.pt
                tf.clear();p=tf.paragraphs[0];p.alignment=PP_ALIGN.CENTER
                for k,v in enumerate(pp):
                    if k:
                        run=p.add_run();run.text=' / ';style_run(run,sz)
                    val=num(v);run=p.add_run();run.text=v.strip()
                    best=val==ranks[k][0];second=len(ranks[k])>1 and val==ranks[k][1]
                    style_run(run,sz,NAVY if best else INK,BOLD if best else FONT,best);run.font.underline=second

set_sections(prs,[('표지와 요약',list(prs.slides)[:3]),(S1,list(prs.slides)[3:7]),(S2,list(prs.slides)[7:12]),(S3,list(prs.slides)[12:18]),(S4,list(prs.slides)[18:23]),(S5,list(prs.slides)[23:26]),(SA,list(prs.slides)[26:])])
fix_table_ids(prs)
prs.core_properties.title='SCARF · Outline v17.1 · v16.3 구성 유지'
scratch=BUILD/'SCARF_v17.1_structure_restored.pptx'
prs.save(scratch)
manifest={'source':str(SRC),'source_sha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),'outline_sha256':hashlib.sha256(MD.read_bytes()).hexdigest(),'backup':str(bak),'build':str(scratch),'slides':35,'main':26,'appendix':9,'policy':'v16.3 slide topics and body roles retained, linear0125 evidence only; unmeasured results explicitly missing'}
(BUILD/'structure_provenance.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print(scratch)
